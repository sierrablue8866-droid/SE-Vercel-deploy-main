# -*- coding: utf-8 -*-
"""
Exports harvested WhatsApp Owners JSON to styled Excel sheet: Owners_Inventory.xlsx
Orders units with photos first, flags units without photos for revision.
"""

import os
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = r"H:\Sheets"
JSON_FILE = os.path.join(OUTPUT_DIR, "Owners_Inventory.json")
EXCEL_FILE = os.path.join(OUTPUT_DIR, "Owners_Inventory.xlsx")

def export_owners_sheet():
    if not os.path.exists(JSON_FILE):
        print(f"No JSON cache found at {JSON_FILE}")
        return

    with open(JSON_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    if not data:
        print("No items in Owners Inventory.")
        return

    df = pd.DataFrame(data)

    # Sort: Units with photos (Has_Photos == 'YES') FIRST, then by Received_At desc
    df["has_photo_rank"] = df["Has_Photos"].map({"YES": 0, "NO": 1})
    df = df.sort_values(["has_photo_rank", "Received_At"], ascending=[True, False]).drop(columns=["has_photo_rank"])

    with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Owners_Inventory", index=False)

        # Summary sheet
        summary_data = {
            "Metric": [
                "Total Owner Units Harvested",
                "Units With Photos (Ready)",
                "Units Without Photos (Needs Revision)",
                "Rent Deals",
                "Sale Deals",
            ],
            "Count": [
                len(df),
                (df["Has_Photos"] == "YES").sum(),
                (df["Has_Photos"] == "NO").sum(),
                (df["Deal_Type"] == "Rent").sum(),
                (df["Deal_Type"] == "Sale").sum(),
            ]
        }
        pd.DataFrame(summary_data).to_excel(writer, sheet_name="Summary", index=False)

    # Apply styling
    wb = openpyxl.load_workbook(EXCEL_FILE)
    hf = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    hfont = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    dfont = Font(name="Segoe UI", size=9)
    photo_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    revision_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

        for c in range(1, ws.max_column + 1):
            ws.cell(1, c).fill = hf
            ws.cell(1, c).font = hfont
            ws.cell(1, c).alignment = Alignment(horizontal="center", vertical="center")

            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, c)
                cell.font = dfont

                # Format price
                if isinstance(cell.value, (int, float)) and "price" in str(ws.cell(1, c).value or "").lower():
                    cell.number_format = "#,##0"

                # Highlight Status column
                if str(ws.cell(1, c).value) == "Status":
                    if "Ready" in str(cell.value or ""):
                        cell.fill = photo_fill
                    elif "Revision" in str(cell.value or ""):
                        cell.fill = revision_fill

            mx = max((len(str(ws.cell(r, c).value or "")) for r in range(1, min(ws.max_row, 100) + 1)), default=10)
            ws.column_dimensions[get_column_letter(c)].width = min(mx + 4, 45)

    wb.save(EXCEL_FILE)
    print(f"✅ Successfully exported Owners Inventory to: {EXCEL_FILE}")

if __name__ == "__main__":
    export_owners_sheet()
