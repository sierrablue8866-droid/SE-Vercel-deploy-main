"""Sierra Estates — Single-sheet inventory consolidator.

Reads the canonical merged workbook (Rental Master + Sales Excluded) and
produces a styled, operational single-sheet Excel file for the ops team.
"""

import re
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

SOURCE = Path('/home/ubuntu/whatsapp_inventory/merged_inventory/merged_rental_inventory.xlsx')
OUT = Path('/home/ubuntu/whatsapp_inventory/merged_inventory/Sierra_Estates_All_Inventory_One_Sheet.xlsx')


def read_formatted_sheet(worksheet):
    """Read a Sierra-formatted openpyxl worksheet and return normalised records."""
    hdr_row = None
    headers = []
    for row in worksheet.iter_rows(
        min_row=1, max_row=min(worksheet.max_row, 20), values_only=True
    ):
        values = [str(v).strip() if v is not None else '' for v in row]
        if 'RecordID' in values or 'UnitFingerprint' in values:
            hdr_row = row
            headers = values
            break
    if hdr_row is None:
        return []

    header_map = {}
    for col_idx, value in enumerate(headers):
        if value and value not in header_map:
            header_map[value] = col_idx

    # Locate the 1-based row index of the header inside the sheet.
    header_index = next(
        i
        for i, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=20, values_only=True),
            start=1,
        )
        if list(row) == list(hdr_row)
    )

    records = []
    for row in worksheet.iter_rows(min_row=header_index + 1, values_only=True):
        if not any(v not in (None, '') for v in row):
            continue
        row_data = {}
        for key, col_idx in header_map.items():
            row_data[key] = row[col_idx] if col_idx < len(row) else None
        if row_data.get('RecordID') or row_data.get('UnitFingerprint'):
            records.append(row_data)
    return records


# ---------------------------------------------------------------------------
# Load source workbook
# ---------------------------------------------------------------------------
src = load_workbook(SOURCE, read_only=True, data_only=True)
rentals = read_formatted_sheet(src['Rental Master'])
sales = read_formatted_sheet(src['Sales Excluded'])

# Keep a single operational inventory sheet while retaining explicit category
# separation via the ListingCategory column.
FIELDS = [
    'RecordID', 'ListingCategory', 'InventoryStatus', 'SourceType',
    'Location', 'Zone', 'PropertyType', 'Bedrooms', 'Bathrooms', 'AreaSqm',
    'PriceEGP', 'Furnished', 'Garden', 'Pool', 'AdditionalFeatures',
    'ContactName', 'ContactPhone', 'OwnerBroker', 'Availability',
    'FollowUpStatus', 'Priority', 'LastContactDate', 'NextAction',
    'AssignedAgent', 'SourceFile', 'SourceSheet', 'SourceRow',
    'UnitFingerprint', 'ConflictFlag', 'ConflictSummary', 'Comment',
    'UpdatedAt',
]

# ---------------------------------------------------------------------------
# Create output workbook
# ---------------------------------------------------------------------------
out = Workbook()
out_ws = out.active
out_ws.title = 'All Inventory Units'
out_ws.sheet_view.showGridLines = False
out_ws.freeze_panes = 'A5'
out_ws.sheet_properties.pageSetUpPr.fitToPage = True
out_ws.page_setup.fitToWidth = 1
out_ws.page_setup.fitToHeight = 0

# Title and context rows
out_ws['A1'] = 'SIERRA ESTATES / ALL INVENTORY UNITS'
out_ws['A1'].font = Font(name='Georgia', size=18, bold=True, color='173B4D')
out_ws['A2'] = (
    'Single-sheet operational inventory consolidated from the canonical '
    'merged workbook. Rental and sale records remain distinguishable by '
    'ListingCategory.'
)
out_ws['A2'].font = Font(name='Calibri', size=10, italic=True, color='5D6B73')
out_ws['A3'] = (
    f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")} '
    f'| Rental Master: {len(rentals):,} | Sales Excluded: {len(sales):,}'
)
out_ws['A3'].font = Font(name='Calibri', size=10, color='5D6B73')

HEADER_ROW = 5
for col_num, field in enumerate(FIELDS, start=1):
    cell = out_ws.cell(HEADER_ROW, col_num, field)
    cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='173B4D')
    cell.alignment = Alignment(
        horizontal='center', vertical='center', wrap_text=True,
    )


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------

def val(rec, *keys, default=None):
    """Return the first non-empty value found in *rec* for the given *keys*."""
    for key in keys:
        if key in rec and rec[key] not in (None, ''):
            return rec[key]
    return default


def normalize(raw_rec, category):
    """Map a source record to the canonical FIELDS list."""
    raw_text = ' '.join(
        str(val(raw_rec, key, default='') or '')
        for key in [
            'RawListingText', 'SourceNotes', 'SourceGroup',
            'Location', 'PropertyType',
        ]
    )
    is_request = (
        category == 'Rental'
        and bool(re.search(
            r'(?i)(mloob|mloba|مطلوب|محتاج|عايز|دورة مطلوب|عميل طالب'
            r'|client looking|looking for|tenant looking|wanted)',
            raw_text,
        ))
    )
    if category == 'Sale':
        listing_category = 'Sale Excluded'
    elif is_request:
        listing_category = 'Client Request'
    else:
        listing_category = 'Available Listing'

    if category == 'Sale':
        status = 'Excluded from rental map'
    else:
        status = 'Pending'

    features = (
        val(raw_rec, 'AdditionalFeatures', 'RawListingText',
            'SourceNotes', default='') or ''
    )
    garden = val(raw_rec, 'Garden', default=False)
    pool = val(raw_rec, 'Pool', default=False)

    if isinstance(garden, str):
        garden = garden.strip().lower() in {
            'true', 'yes', 'y', '1', 'garden', 'حديقة', 'جاردن',
        }
    if isinstance(pool, str):
        pool = pool.strip().lower() in {
            'true', 'yes', 'y', '1', 'pool', 'swimming pool',
            'حمام سباحة', 'بول',
        }
    if not garden:
        garden = bool(re.search(
            r'(?i)(garden|حديقة|جاردن|private garden)', raw_text,
        ))
    if not pool:
        pool = bool(re.search(
            r'(?i)(pool|حمام سباحة|سباحة|private pool)', raw_text,
        ))

    return {
        'RecordID': val(raw_rec, 'RecordID', 'Code'),
        'ListingCategory': listing_category,
        'InventoryStatus': status,
        'SourceType': val(raw_rec, 'SourceType', 'Source', default='Unknown'),
        'Location': val(raw_rec, 'Location', 'Compound', 'Area'),
        'Zone': val(raw_rec, 'Zone'),
        'PropertyType': val(raw_rec, 'PropertyType', 'Type'),
        'Bedrooms': val(raw_rec, 'Bedrooms', 'Beds', 'bedrooms'),
        'Bathrooms': val(raw_rec, 'Bathrooms', 'Baths'),
        'AreaSqm': val(raw_rec, 'AreaSqm', 'Area', 'Space'),
        'PriceEGP': val(
            raw_rec, 'PriceEGP', 'MonthlyRentEGP', 'Price',
            'UnitPrice', 'Rent',
        ),
        'Furnished': val(
            raw_rec, 'Furnished', 'Furnishing', 'FurnishedStatus',
        ),
        'Garden': bool(garden),
        'Pool': bool(pool),
        'AdditionalFeatures': features,
        'ContactName': val(raw_rec, 'ContactName', 'SourceName', 'Name'),
        'ContactPhone': val(
            raw_rec, 'ContactPhone', 'SourceContact', 'Phone', 'Mobile',
        ),
        'OwnerBroker': val(
            raw_rec, 'OwnerBroker', 'Ownership', 'SourceType', 'Owner',
        ),
        'Availability': val(raw_rec, 'Availability', 'Availablty'),
        'FollowUpStatus': val(
            raw_rec, 'FollowUpStatus',
            default='Pending' if category == 'Rental' else 'Not Applicable',
        ),
        'Priority': val(raw_rec, 'Priority', default='Normal'),
        'LastContactDate': val(raw_rec, 'LastContactDate'),
        'NextAction': val(raw_rec, 'NextAction'),
        'AssignedAgent': val(
            raw_rec, 'AssignedAgent', default='Unassigned',
        ),
        'SourceFile': val(raw_rec, 'SourceFile', 'SourceWorkbook'),
        'SourceSheet': val(raw_rec, 'SourceSheet'),
        'SourceRow': val(raw_rec, 'SourceRow'),
        'UnitFingerprint': val(
            raw_rec, 'UnitFingerprint', 'DedupGroupKey',
        ),
        'ConflictFlag': val(raw_rec, 'ConflictFlag', default=False),
        'ConflictSummary': val(
            raw_rec, 'ConflictSummary', 'ConflictNotes',
        ),
        'Comment': val(raw_rec, 'Comment', 'SourceNotes'),
        'UpdatedAt': val(raw_rec, 'UpdatedAt', 'ListingDate'),
    }


# ---------------------------------------------------------------------------
# Populate rows
# ---------------------------------------------------------------------------
rows = (
    [normalize(r, 'Rental') for r in rentals]
    + [normalize(r, 'Sale') for r in sales]
)
for row_idx, unit in enumerate(rows, start=HEADER_ROW + 1):
    for col_num, field in enumerate(FIELDS, start=1):
        cell = out_ws.cell(row_idx, col_num, unit.get(field))
        cell.font = Font(name='Calibri', size=10, color='1F2D33')
        cell.alignment = Alignment(
            vertical='top',
            wrap_text=field in {
                'AdditionalFeatures', 'ConflictSummary', 'Comment',
            },
        )
        if field == 'PriceEGP' and isinstance(cell.value, (int, float)):
            cell.number_format = '#,##0'
        if field in {'Garden', 'Pool', 'ConflictFlag'}:
            cell.alignment = Alignment(horizontal='center', vertical='top')
    if row_idx % 2 == 0:
        for c in range(1, len(FIELDS) + 1):
            out_ws.cell(row_idx, c).fill = PatternFill(
                'solid', fgColor='F4F7F7',
            )

# ---------------------------------------------------------------------------
# Table and formatting
# ---------------------------------------------------------------------------
LAST_ROW = HEADER_ROW + len(rows)
LAST_COL = len(FIELDS)
TABLE_REF = (
    f'A{HEADER_ROW}:'
    f'{chr(64 + LAST_COL) if LAST_COL <= 26 else "AF"}{LAST_ROW}'
)
tab = Table(displayName='AllInventoryUnits', ref=TABLE_REF)
tab.tableStyleInfo = TableStyleInfo(
    name='TableStyleMedium2',
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
out_ws.add_table(tab)

# Conditional formatting for operational fields.
STATUS_COL = FIELDS.index('FollowUpStatus') + 1
PRIORITY_COL = FIELDS.index('Priority') + 1
CATEGORY_COL = FIELDS.index('ListingCategory') + 1
for col_num in [STATUS_COL, PRIORITY_COL, CATEGORY_COL]:
    letter = out_ws.cell(HEADER_ROW, col_num).column_letter
    out_ws.conditional_formatting.add(
        f'{letter}{HEADER_ROW + 1}:{letter}{LAST_ROW}',
        CellIsRule(
            operator='equal',
            formula=['"Pending"'],
            fill=PatternFill('solid', fgColor='FFF3CD'),
        ),
    )

# Column widths
WIDTHS = {
    'RecordID': 17, 'ListingCategory': 19, 'InventoryStatus': 22,
    'SourceType': 13, 'Location': 22, 'Zone': 15, 'PropertyType': 18,
    'Bedrooms': 10, 'Bathrooms': 10, 'AreaSqm': 11, 'PriceEGP': 13,
    'Furnished': 12, 'Garden': 9, 'Pool': 9, 'AdditionalFeatures': 32,
    'ContactName': 22, 'ContactPhone': 17, 'OwnerBroker': 14,
    'Availability': 15, 'FollowUpStatus': 16, 'Priority': 11,
    'LastContactDate': 16, 'NextAction': 28, 'AssignedAgent': 16,
    'SourceFile': 28, 'SourceSheet': 18, 'SourceRow': 10,
    'UnitFingerprint': 28, 'ConflictFlag': 12, 'ConflictSummary': 30,
    'Comment': 34, 'UpdatedAt': 18,
}
for ci, field in enumerate(FIELDS, start=1):
    col_letter = out_ws.cell(HEADER_ROW, ci).column_letter
    out_ws.column_dimensions[col_letter].width = WIDTHS.get(field, 14)
out_ws.row_dimensions[1].height = 28
out_ws.row_dimensions[2].height = 30
out_ws.row_dimensions[HEADER_ROW].height = 38
out_ws.auto_filter.ref = TABLE_REF
out_ws.print_title_rows = f'{HEADER_ROW}:{HEADER_ROW}'

# ---------------------------------------------------------------------------
# Write output
# ---------------------------------------------------------------------------
OUT.parent.mkdir(parents=True, exist_ok=True)
out.save(OUT)
print(f'Wrote {OUT}')
print(f'Rows: {len(rows)} | rental: {len(rentals)} | sale: {len(sales)}')
