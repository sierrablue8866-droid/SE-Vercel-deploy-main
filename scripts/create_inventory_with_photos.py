from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path('/home/ubuntu/whatsapp_inventory/merged_inventory/Sierra_Estates_All_Inventory_One_Sheet.xlsx')
OUT = Path('/home/ubuntu/SE-Vercel-deploy-main/Inventory_with_Photos.xlsx')

# The inspected WhatsApp archives contain no image/media files. Do not add
# fabricated photo rows, placeholder URLs, or a misleading photo status column.
# Keep the validated inventory schema unchanged and reserve this filename for the
# next media-inclusive WhatsApp export.
wb = load_workbook(SOURCE)
ws = wb['All Inventory Units']
ws['A1'] = 'SIERRA ESTATES / INVENTORY WITH PHOTOS'
ws['A2'] = 'Unified inventory reserved for photo enrichment. No WhatsApp media was supplied in the source archives, so no photo fields or placeholders were added.'
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f'Wrote {OUT}')
print(f'Rows: {ws.max_row - 5}')
print('Photo media: none supplied; workbook schema preserved without placeholders')
