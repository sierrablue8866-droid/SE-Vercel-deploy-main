#!/usr/bin/env python3
"""Explore all inventory source files: headers + value distributions."""
import openpyxl
from collections import Counter

XLSX = '/home/z/my-project/sierra-blu/data/Sierra_Estates_Consolidated_Master_Inventory_All_Sheets.xlsx'

wb = openpyxl.load_workbook(XLSX, read_only=True)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n{'='*80}\nSHEET: {name}  ({ws.max_row} rows x {ws.max_column} cols)\n{'='*80}")
    rows = ws.iter_rows(min_row=1, max_row=3, values_only=True)
    headers = None
    for i, r in enumerate(rows):
        if i == 0:
            headers = list(r)
            print("HEADERS:", headers)
        else:
            print(f"ROW{i}:", [str(x)[:40] if x is not None else None for x in r])
wb.close()

# Sample value distributions for All Master Listings key columns
wb = openpyxl.load_workbook(XLSX, read_only=True)
ws = wb['All Master Listings']
compounds = Counter(); deals = Counter(); segs = Counter(); statuses = Counter(); ptypes = Counter()
n = 0
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
    if r[0] is None and r[3] is None: continue
    n += 1
    segs[r[1]] += 1
    deals[r[2]] += 1
    if r[3]: compounds[str(r[3]).strip()] += 1
    if r[13]: statuses[str(r[13]).strip()] += 1
    if r[4]: ptypes[str(r[4]).strip()] += 1
wb.close()

print(f"\n\nALL MASTER LISTINGS: {n} data rows")
print("\nSEGMENTS:", dict(segs))
print("\nDEAL TYPES:", dict(deals))
print("\nSTATUS VALUES (top 25):", statuses.most_common(25))
print("\nPROPERTY TYPES (top 30):", ptypes.most_common(30))
print(f"\nUNIQUE COMPOUNDS: {len(compounds)}")
print("TOP 40:", compounds.most_common(40))
