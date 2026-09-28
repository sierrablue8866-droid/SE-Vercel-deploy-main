#!/usr/bin/env python3
import sys, io, warnings, re
from datetime import datetime
from pathlib import Path

warnings.filterwarnings("ignore")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(r"H:\last\Main\SE-Vercel-deploy-main")
SOURCE    = Path(r"C:\Users\Sierr\Downloads\20-7\20-7-2026.xlsx")
MASTER    = REPO_ROOT / "data" / "Sierra_Estates_Rent_Master_Inventory.xlsx"
OUT_DATE  = datetime.now().strftime("%Y-%m-%d")
OUT_PATH  = REPO_ROOT / "data" / f"Sierra_Estates_Master_Combined_MERGED_{OUT_DATE}.xlsx"

print(f"Source : {SOURCE}")
print(f"Master : {MASTER}")
print(f"Output : {OUT_PATH}")

def norm_phone(v):
    if not v: return ""
    s = re.sub(r"[^\d]", "", str(v))
    if s.startswith("002"): s = s[3:]
    elif s.startswith("20") and len(s) > 10: s = s[2:]
    return s.lstrip("0")[-9:]

def norm_text(v):
    if not v: return ""
    return str(v).strip().lower()

def enrich(existing, new):
    result = dict(existing)
    for k, v in new.items():
        if k not in result or result[k] in (None, "", "None"):
            if v not in (None, "", "None"):
                result[k] = v
        if k in ("price","bedrooms","bathrooms","area","availability","furnishing"):
            if v not in (None, "", "None", 0):
                result[k] = v
    return result

def sheet_to_dicts(wb, sheet_name):
    if sheet_name not in wb.sheetnames: return [], []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return [], []
    start = 0
    while start < len(rows) and all(v is None for v in rows[start]):
        start += 1
    if start >= len(rows): return [], []
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[start])]
    data = []
    for row in rows[start+1:]:
        if all(v is None for v in row): continue
        d = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        data.append(d)
    return headers, data

print("\n-- Loading workbooks --")
src_wb = openpyxl.load_workbook(SOURCE, data_only=True)
mst_wb = openpyxl.load_workbook(MASTER, data_only=True)

# ---- Parse SOURCE Owners-Rent ----
_, src_or_rows = sheet_to_dicts(src_wb, "Owners-Rent")
print(f"Owners-Rent rows: {len(src_or_rows)}")

def parse_owners_rent_src(rows):
    out = []
    for r in rows:
        phone = norm_phone(r.get("Mobile") or r.get("Mobile ") or "")
        compound = norm_text(r.get("Location ") or r.get("Location") or "")
        if not phone: continue
        out.append({
            "phone": phone, "name": str(r.get("Name") or r.get("Owner") or "").strip(),
            "compound": compound, "code": str(r.get("Code") or "").strip(),
            "availability": str(r.get("Availablty") or "").strip(),
            "bedrooms": r.get("bedrooms") or "", "price": r.get("Unit Price") or "",
            "furnishing": str(r.get("Furnished or not") or "").strip(),
            "property_type": str(r.get("Property Tybe") or r.get("Type") or "").strip(),
            "area": r.get("Area") or "", "garden": str(r.get("Garden") or "").strip(),
            "timestamp": str(r.get("Timestamp") or "").strip(),
            "operation": "rent", "source": "20-7-2026",
        })
    return out

src_or = parse_owners_rent_src(src_or_rows)
print(f"Owners-Rent parsed: {len(src_or)}")

# ---- Parse SOURCE Owners-Resale ----
sheet_name_ors = "Owners-Resale " if "Owners-Resale " in src_wb.sheetnames else "Owners-Resale"
_, src_ors_rows = sheet_to_dicts(src_wb, sheet_name_ors)
print(f"Owners-Resale rows: {len(src_ors_rows)}")

def parse_owners_resale_src(rows):
    out = []
    for r in rows:
        phone = norm_phone(r.get("تليفون") or r.get("Mobile") or "")
        compound = norm_text(r.get("الكمبوند") or r.get("Compound") or "")
        if not phone: continue
        out.append({
            "phone": phone, "name": str(r.get("Name") or "").strip(),
            "compound": compound, "code": str(r.get("الكود") or r.get("Code") or "").strip(),
            "availability": str(r.get("متاحه /غير متاحه") or "").strip(),
            "price": r.get("السعر") or "",
            "property_type": str(r.get("نوع الوحده") or "").strip(),
            "area": r.get("المساحه") or "", "garden": str(r.get("الحديقة") or "").strip(),
            "finishing": str(r.get("التشطيب") or "").strip(),
            "description": str(r.get("بيان الوحده") or "").strip(),
            "bedrooms": r.get("الغرف") or "", "bathrooms": r.get("الحمامات") or "",
            "operation": "resale", "source": "20-7-2026",
        })
    return out

src_ors = parse_owners_resale_src(src_ors_rows)
print(f"Owners-Resale parsed: {len(src_ors)}")

# ---- Parse SOURCE Brokers Rent (no header row) ----
ws_br = src_wb["Brokers Rent"]
br_all = list(ws_br.iter_rows(values_only=True))
br_hdrs = ["Timestamp","Role","Name","Mobile","Code","Location","Price","PaymentType","ContactName","ContactRole","Furnishing","PropertyType","Operation"]

def parse_brokers_rent_src(all_rows, hdrs):
    out = []
    for row in all_rows:
        if all(v is None for v in row): continue
        d = {hdrs[i] if i < len(hdrs) else f"col_{i}": row[i] for i in range(len(row))}
        phone = norm_phone(d.get("Mobile") or "")
        compound = norm_text(d.get("Location") or "")
        if not phone: continue
        out.append({
            "phone": phone, "name": str(d.get("Name") or "").strip(),
            "compound": compound, "code": str(d.get("Code") or "").strip(),
            "price": d.get("Price") or "",
            "furnishing": str(d.get("Furnishing") or "").strip(),
            "property_type": str(d.get("PropertyType") or "").strip(),
            "operation": str(d.get("Operation") or "rent").strip().lower(),
            "payment_type": str(d.get("PaymentType") or "").strip(),
            "role": str(d.get("Role") or "broker").strip(),
            "source": "20-7-2026",
        })
    return out

src_br = parse_brokers_rent_src(br_all, br_hdrs)
print(f"Brokers Rent parsed: {len(src_br)}")

# ---- Parse SOURCE Team Units ----
_, src_tu_rows = sheet_to_dicts(src_wb, "Team Units")
print(f"Team Units rows: {len(src_tu_rows)}")

def parse_team_units_src(rows):
    out = []
    for r in rows:
        compound = norm_text(r.get("الموقع (الكمبوند)") or "")
        out.append({
            "bedrooms": r.get("الغرف") or "", "price": r.get("السعر (EGP)") or "",
            "property_type": str(r.get("نوع العقار") or "").strip(),
            "compound": compound, "bathrooms": r.get("الحمامات") or "",
            "area": r.get("المساحة") or "",
            "operation": str(r.get("العملية") or "rent").strip().lower(),
            "date": str(r.get("التاريخ") or "").strip(),
            "sender": str(r.get("المرسل") or "").strip(),
            "group": str(r.get("اسم الجروب") or "").strip(),
            "source": "20-7-2026-TeamUnits",
        })
    return out

src_tu = parse_team_units_src(src_tu_rows)
print(f"Team Units parsed: {len(src_tu)}")

# ---- Parse MASTER Direct Owners Rent ----
_, mst_dor_rows = sheet_to_dicts(mst_wb, "Direct Owners Rent (298)")
print(f"\nMaster Owners Rent rows: {len(mst_dor_rows)}")

def parse_master_or(rows):
    out = []
    for r in rows:
        phone = norm_phone(r.get("Owner Phone") or "")
        compound = norm_text(r.get("Compound / Community") or r.get("Compound") or "")
        if not phone: continue
        out.append({
            "phone": phone, "name": str(r.get("Owner / Contact Name") or "").strip(),
            "compound": compound, "code": str(r.get("Unit Code") or "").strip(),
            "zone": str(r.get("Zone / Area") or "").strip(),
            "property_type": str(r.get("Property Type") or "").strip(),
            "price": r.get("Monthly Rent (EGP)") or "",
            "price_display": str(r.get("Rent Display") or "").strip(),
            "area": r.get("Area (sqm)") or "",
            "bedrooms": r.get("Bedrooms") or "", "bathrooms": r.get("Bathrooms") or "",
            "furnishing": str(r.get("Furnishing") or "").strip(),
            "whatsapp": str(r.get("Direct WhatsApp") or "").strip(),
            "listing_status": str(r.get("Listing Status") or "").strip(),
            "source_channel": str(r.get("Source Channel") or "").strip(),
            "availability": "", "operation": "rent", "source": "master",
        })
    return out

mst_dor = parse_master_or(mst_dor_rows)

# ---- Parse MASTER Broker Rent ----
_, mst_brn_rows = sheet_to_dicts(mst_wb, "Broker Rent Network (4955)")
print(f"Master Broker rows: {len(mst_brn_rows)}")

def parse_master_br(rows):
    out = []
    for r in rows:
        phone = norm_phone(r.get("Contact Phone") or r.get("Phone") or "")
        compound = norm_text(r.get("Compound") or r.get("Location") or "")
        if not phone: continue
        out.append({
            "phone": phone, "record_id": str(r.get("RecordID") or "").strip(),
            "unit_code": str(r.get("UnitCode") or "").strip(),
            "name": str(r.get("Contact Name") or "").strip(),
            "compound": compound, "location": str(r.get("Location") or "").strip(),
            "zone": str(r.get("Zone") or "").strip(),
            "property_type": str(r.get("PropertyType") or "").strip(),
            "operation": str(r.get("Operation") or "rent").strip().lower(),
            "price": r.get("Price (EGP)") or "",
            "price_fmt": str(r.get("Price Formatted") or "").strip(),
            "area": r.get("Area (sqm)") or "",
            "bedrooms": r.get("Bedrooms") or "", "bathrooms": r.get("Bathrooms") or "",
            "furnishing": str(r.get("Furnishing") or "").strip(),
            "payment_type": "", "role": "broker", "source": "master",
        })
    return out

mst_brn = parse_master_br(mst_brn_rows)

# ---- MERGE: Owners Rent ----
print("\n-- Merging Owners Rent --")
mst_or_idx = {(r["phone"], r["compound"]): r for r in mst_dor}
src_or_idx = {(r["phone"], r["compound"]): r for r in src_or}

added = updated = kept = 0
merged_or = {}
for key, mrec in mst_or_idx.items():
    if key in src_or_idx:
        m = enrich(mrec, src_or_idx[key])
        m["source"] = "master+20-7"
        merged_or[key] = m; updated += 1
    else:
        merged_or[key] = mrec; kept += 1
for key, srec in src_or_idx.items():
    if key not in merged_or:
        merged_or[key] = srec; added += 1

print(f"  kept={kept}, updated={updated}, added={added}, total={len(merged_or)}")

# ---- MERGE: Broker Rent ----
print("-- Merging Broker Rent --")
mst_br_idx = {(r["phone"], r.get("compound","")): r for r in mst_brn}
src_br_idx = {(r["phone"], r.get("compound","")): r for r in src_br}

br_added = br_updated = br_kept = 0
merged_br = {}
for key, mrec in mst_br_idx.items():
    if key in src_br_idx:
        m = enrich(mrec, src_br_idx[key])
        m["source"] = "master+20-7"
        merged_br[key] = m; br_updated += 1
    else:
        merged_br[key] = mrec; br_kept += 1
for key, srec in src_br_idx.items():
    if key not in merged_br:
        merged_br[key] = srec; br_added += 1

print(f"  kept={br_kept}, updated={br_updated}, added={br_added}, total={len(merged_br)}")

# ---- DEDUP: Owners Resale ----
print("-- Deduplicating Owners Resale --")
merged_ors = {}
for rec in src_ors:
    key = (rec["phone"], rec["compound"])
    if key not in merged_ors:
        merged_ors[key] = rec
    else:
        ex = merged_ors[key]
        if sum(1 for v in rec.values() if v not in (None,"","None")) > sum(1 for v in ex.values() if v not in (None,"","None")):
            merged_ors[key] = rec
print(f"  raw={len(src_ors)}, deduped={len(merged_ors)}")

# ---- DEDUP: Team Units ----
print("-- Deduplicating Team Units --")
merged_tu = {}
for rec in src_tu:
    key = (rec["compound"], str(rec.get("bedrooms","")), str(rec.get("price","")), rec.get("operation",""))
    if key not in merged_tu:
        merged_tu[key] = rec
print(f"  raw={len(src_tu)}, deduped={len(merged_tu)}")

# ---- WRITE OUTPUT ----
print("\n-- Writing output --")
out_wb = openpyxl.Workbook()
out_wb.remove(out_wb.active)

HFILL = PatternFill("solid", fgColor="1F3864")
HFONT = Font(bold=True, color="FFFFFF", size=11)
THIN  = Side(style="thin", color="CCCCCC")
BRD   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AFILL = PatternFill("solid", fgColor="E2EFDA")
NFILL = PatternFill("solid", fgColor="FCE4D6")
SFILL = {"master": PatternFill("solid", fgColor="FFFFFF"), "20-7-2026": PatternFill("solid", fgColor="EBF3FB"), "master+20-7": PatternFill("solid", fgColor="FFF2CC")}

def whdr(ws, cols, row=1):
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HFILL; cell.font = HFONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BRD

def wrow(ws, ri, vals, fill=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=ri, column=c, value=v)
        cell.border = BRD
        cell.alignment = Alignment(vertical="center")
        if fill: cell.fill = fill

def awidth(ws):
    for col in ws.columns:
        ml = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(ml+2, 10), 42)

# Sheet 1: Summary
ws_s = out_wb.create_sheet("Summary")
ws_s.sheet_properties.tabColor = "1F3864"
whdr(ws_s, ["Segment","Count","Source","Notes"])
rows_sum = [
    ("Direct Owners Rent",  len(merged_or),  "Master+20-7", f"kept={kept}, updated={updated}, added={added}"),
    ("Broker Rent Network", len(merged_br),  "Master+20-7", f"kept={br_kept}, updated={br_updated}, added={br_added}"),
    ("Owners Resale",       len(merged_ors), "20-7-2026",   f"raw={len(src_ors)}, deduped={len(merged_ors)}"),
    ("Team Units",          len(merged_tu),  "20-7-2026",   f"raw={len(src_tu)}, deduped={len(merged_tu)}"),
    ("GRAND TOTAL", len(merged_or)+len(merged_br)+len(merged_ors)+len(merged_tu), "--", datetime.now().strftime("%Y-%m-%d %H:%M")),
]
for i, row in enumerate(rows_sum, 2):
    for c, v in enumerate(row, 1):
        cell = ws_s.cell(row=i, column=c, value=v)
        cell.border = BRD
        if row[0] == "GRAND TOTAL":
            cell.fill = HFILL; cell.font = Font(bold=True, color="FFFFFF")
awidth(ws_s)

# Sheet 2: Owners Rent
or_cols = ["Unit Code","Name","Phone","Compound","Zone","Property Type","Bedrooms","Bathrooms","Area (sqm)","Monthly Rent (EGP)","Furnishing","Availability","Garden","WhatsApp","Listing Status","Source","Operation"]
ws_or = out_wb.create_sheet(f"Owners Rent ({len(merged_or)})")
ws_or.sheet_properties.tabColor = "00B050"
whdr(ws_or, or_cols)
ws_or.freeze_panes = "A2"
ws_or.auto_filter.ref = f"A1:{get_column_letter(len(or_cols))}1"
for i, rec in enumerate(sorted(merged_or.values(), key=lambda r: (r.get("compound",""), r.get("name",""))), 2):
    av = rec.get("availability","")
    sf = "available" in str(av).lower()
    sk = "master+20-7" if "master" in rec.get("source","") and "20-7" in rec.get("source","") else rec.get("source","master")
    fill = AFILL if sf else (SFILL.get(sk, SFILL["master"]) if not av else NFILL)
    wrow(ws_or, i, [rec.get("code",""), rec.get("name",""), rec.get("phone",""), rec.get("compound",""), rec.get("zone",""), rec.get("property_type",""), rec.get("bedrooms",""), rec.get("bathrooms",""), rec.get("area",""), rec.get("price",""), rec.get("furnishing",""), rec.get("availability",""), rec.get("garden",""), rec.get("whatsapp",""), rec.get("listing_status",""), rec.get("source",""), rec.get("operation","rent")], fill)
awidth(ws_or)

# Sheet 3: Broker Rent
br_cols = ["Record ID","Unit Code","Phone","Name","Compound","Location","Zone","Property Type","Operation","Price (EGP)","Price Formatted","Area (sqm)","Bedrooms","Bathrooms","Furnishing","Payment Type","Role","Source"]
ws_br2 = out_wb.create_sheet(f"Broker Rent ({len(merged_br)})")
ws_br2.sheet_properties.tabColor = "0070C0"
whdr(ws_br2, br_cols)
ws_br2.freeze_panes = "A2"
ws_br2.auto_filter.ref = f"A1:{get_column_letter(len(br_cols))}1"
for i, rec in enumerate(sorted(merged_br.values(), key=lambda r: (r.get("compound",""), r.get("name",""))), 2):
    sk = rec.get("source","master")
    fill = SFILL.get(sk, SFILL["master"])
    wrow(ws_br2, i, [rec.get("record_id",""), rec.get("unit_code","") or rec.get("code",""), rec.get("phone",""), rec.get("name",""), rec.get("compound",""), rec.get("location","") or rec.get("compound",""), rec.get("zone",""), rec.get("property_type",""), rec.get("operation","rent"), rec.get("price",""), rec.get("price_fmt",""), rec.get("area",""), rec.get("bedrooms",""), rec.get("bathrooms",""), rec.get("furnishing",""), rec.get("payment_type",""), rec.get("role","broker"), rec.get("source","")], fill)
awidth(ws_br2)

# Sheet 4: Owners Resale
ors_cols = ["Code","Name","Phone","Compound","Property Type","Price (EGP)","Area (sqm)","Bedrooms","Bathrooms","Garden","Finishing","Availability","Description","Source","Operation"]
ws_ors2 = out_wb.create_sheet(f"Owners Resale ({len(merged_ors)})")
ws_ors2.sheet_properties.tabColor = "FF6600"
whdr(ws_ors2, ors_cols)
ws_ors2.freeze_panes = "A2"
ws_ors2.auto_filter.ref = f"A1:{get_column_letter(len(ors_cols))}1"
for i, rec in enumerate(sorted(merged_ors.values(), key=lambda r: (r.get("compound",""), r.get("name",""))), 2):
    av = str(rec.get("availability",""))
    fill = AFILL if "متاحه" in av and "غير" not in av else NFILL
    wrow(ws_ors2, i, [rec.get("code",""), rec.get("name",""), rec.get("phone",""), rec.get("compound",""), rec.get("property_type",""), rec.get("price",""), rec.get("area",""), rec.get("bedrooms",""), rec.get("bathrooms",""), rec.get("garden",""), rec.get("finishing",""), rec.get("availability",""), rec.get("description",""), rec.get("source",""), rec.get("operation","resale")], fill)
awidth(ws_ors2)

# Sheet 5: Team Units
tu_cols = ["Compound","Property Type","Operation","Bedrooms","Bathrooms","Area (sqm)","Price (EGP)","Date","Sender","Group","Source"]
ws_tu2 = out_wb.create_sheet(f"Team Units ({len(merged_tu)})")
ws_tu2.sheet_properties.tabColor = "7030A0"
whdr(ws_tu2, tu_cols)
ws_tu2.freeze_panes = "A2"
ws_tu2.auto_filter.ref = f"A1:{get_column_letter(len(tu_cols))}1"
for i, rec in enumerate(sorted(merged_tu.values(), key=lambda r: (r.get("compound",""), r.get("operation",""))), 2):
    op = str(rec.get("operation","")).lower()
    fill = AFILL if "rent" in op else (NFILL if "sale" in op else None)
    wrow(ws_tu2, i, [rec.get("compound",""), rec.get("property_type",""), rec.get("operation",""), rec.get("bedrooms",""), rec.get("bathrooms",""), rec.get("area",""), rec.get("price",""), rec.get("date",""), rec.get("sender",""), rec.get("group",""), rec.get("source","")], fill)
awidth(ws_tu2)

out_wb.save(OUT_PATH)
print(f"\n✅ Saved: {OUT_PATH}")
print(f"   Owners Rent   : {len(merged_or)}")
print(f"   Broker Rent   : {len(merged_br)}")
print(f"   Owners Resale : {len(merged_ors)}")
print(f"   Team Units    : {len(merged_tu)}")
print(f"   GRAND TOTAL   : {len(merged_or)+len(merged_br)+len(merged_ors)+len(merged_tu)}")
print("Done.")
