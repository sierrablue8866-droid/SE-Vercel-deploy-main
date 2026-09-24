#!/usr/bin/env python3
"""
Sierra Estates - Master Real Estate Inventory Consolidator & Deduplicator
========================================================================
Merges all historical and active datasets:
1. Inventory_with_Photos.xlsx (7,655 units with media links)
2. sierra-estates-master-inventory.xlsx (16,079 units from scrapers/archives)
3. Sierra_Estates_Owners_Units_Rent_and_Resale.xlsx (585 verified owner units)
4. inventory_master_unified.csv (8,098 units)

Features:
- Strict multi-field deduplication (Phone last 7 digits, Price, Compound, Deal Type).
- Resolves duplicates by merging fields (keeps Owner status over Broker, merges photo URLs, keeps longest description).
- Clear categorization: Owner vs Broker, Rent vs Sale/Re-sale.
- Single master Excel file with 6 organized sheets:
  1. Summary (KPI dashboard & statistics)
  2. All_Units (Master catalog)
  3. Owners_Rent (Direct owners rentals)
  4. Owners_Sale (Direct owners resale/sale)
  5. Brokers_Rent (Broker network rentals)
  6. Brokers_Sale (Broker network resale/sale)
- High-end OpenPyXL styling (branded colors, auto-filter, frozen panes, numeric formats).
"""

import os
import re
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def clean_phone(val):
    if not val:
        return ''
    s = re.sub(r'[^0-9]', '', str(val))
    if s.startswith('20') and len(s) == 12:
        s = s[2:]
    elif s.startswith('0020') and len(s) == 14:
        s = s[4:]
    if len(s) == 10 and s.startswith('1'):
        s = '0' + s
    return s if (len(s) == 11 and s.startswith('01')) else s

def clean_price(val):
    if val is None or val == '':
        return 0
    try:
        if isinstance(val, (int, float)):
            return float(val)
        cleaned = re.sub(r'[^0-9.]', '', str(val))
        return float(cleaned) if cleaned else 0
    except:
        return 0

def clean_num(val):
    if val is None or val == '':
        return 0
    try:
        if isinstance(val, (int, float)):
            return int(val)
        cleaned = re.sub(r'[^0-9]', '', str(val))
        return int(cleaned) if cleaned else 0
    except:
        return 0

def normalize_compound(text):
    t = str(text or '').strip().lower()
    if not t or t in ('none', 'null', 'nan', '-', ''):
        return 'New Cairo'
    if 'mivida' in t or 'ميفيدا' in t or 'mevida' in t:
        return 'Mivida (Emaar)'
    if 'fifth square' in t or 'فFifth square' in t or 'المراسم' in t or 'marasem' in t:
        return 'Fifth Square (Al Marasem)'
    if 'rehab' in t or 'الرحاب' in t:
        return 'Al Rehab'
    if 'madinaty' in t or 'مدينتي' in t:
        return 'Madinaty'
    if 'hyde park' in t or 'هايد بارك' in t:
        return 'Hyde Park'
    if 'palm hills' in t or 'بالم هيلز' in t:
        return 'Palm Hills New Cairo'
    if 'mountain view' in t or 'ماونتن فيو' in t:
        return 'Mountain View iCity'
    if 'sodic' in t or 'سوديك' in t or 'villette' in t or 'فيليت' in t or 'east town' in t or 'easttown' in t:
        return 'Villette by SODIC'
    if 'cairo plaza' in t or 'كايرو بلازا' in t:
        return 'Cairo Plaza Towers'
    if 'uptown' in t or 'اب تاون' in t or 'أب تاون' in t:
        return 'Uptown Cairo'
    if 'waterway' in t or 'وتر واي' in t or 'واتر واي' in t:
        return 'The Waterway'
    if 'swan lake' in t or 'سوان ليك' in t:
        return 'Swan Lake Residences'
    if 'katameya' in t or 'قطامية' in t or 'dunes' in t:
        return 'Katameya Dunes / Heights'
    if 'lake view' in t or 'ليك فيو' in t:
        return 'Lake View'
    if 'beit el watan' in t or 'بيت الوطن' in t:
        return 'Beit El Watan'
    if 'tagamoa' in t or 'تجمع' in t or 'new cairo' in t:
        return 'New Cairo Prime'
    return str(text).strip()

def normalize_deal(text):
    t = str(text or '').strip().lower()
    if 'rent' in t or 'إيجار' in t or 'ايجار' in t or 'mo' in t:
        return 'Rent'
    return 'Sale / Re-sale'

def normalize_party(party_text, channel_text=''):
    p = str(party_text or '').strip().lower()
    c = str(channel_text or '').strip().lower()
    if 'owner' in p or 'مالك' in p or 'direct' in p or 'owner' in c or 'مالك' in c:
        return 'Owner'
    return 'Broker'

def build_dedup_key(unit):
    phone = unit['phone']
    deal = unit['deal_type'].lower()
    price_rounded = int(round(unit['price_egp'], -3)) if unit['price_egp'] > 0 else 0
    compound = unit['compound'].lower()[:8]
    code = unit['unit_code'].upper().strip()

    if len(phone) >= 7:
        return f"p:{phone[-7:]}|d:{deal}|pr:{price_rounded}|c:{compound}"
    elif len(code) >= 4 and not code.startswith('MS-') and not code.startswith('SE-'):
        return f"c:{code}|d:{deal}|c:{compound}"
    else:
        area = int(unit['area_sqm']) if unit['area_sqm'] > 0 else 0
        beds = int(unit['bedrooms']) if unit['bedrooms'] > 0 else 0
        return f"s:{deal}|pr:{price_rounded}|a:{area}|b:{beds}|c:{compound}"

def merge_units(existing, incoming):
    existing['duplicates_count'] = existing.get('duplicates_count', 1) + 1
    
    # Priority: Owner over Broker
    if incoming['advertiser_type'] == 'Owner':
        existing['advertiser_type'] = 'Owner'
        
    # Prefer non-empty phone
    if not existing['phone'] and incoming['phone']:
        existing['phone'] = incoming['phone']
        existing['whatsapp_link'] = incoming['whatsapp_link']
        
    # Prefer non-empty name
    if not existing['contact_name'] and incoming['contact_name']:
        existing['contact_name'] = incoming['contact_name']
        
    # Merge photo URLs
    if incoming['photo_urls']:
        if existing['photo_urls']:
            urls = set(existing['photo_urls'].split(', ') + incoming['photo_urls'].split(', '))
            existing['photo_urls'] = ', '.join([u.strip() for u in urls if u.strip()])
        else:
            existing['photo_urls'] = incoming['photo_urls']
        existing['has_photos'] = 'Yes'
        
    # Keep longest description
    if len(incoming['description']) > len(existing['description']):
        existing['description'] = incoming['description']
        
    # Keep non-zero specs
    if existing['area_sqm'] == 0 and incoming['area_sqm'] > 0:
        existing['area_sqm'] = incoming['area_sqm']
    if existing['bedrooms'] == 0 and incoming['bedrooms'] > 0:
        existing['bedrooms'] = incoming['bedrooms']
    if existing['bathrooms'] == 0 and incoming['bathrooms'] > 0:
        existing['bathrooms'] = incoming['bathrooms']

def main():
    print("==================================================")
    print("SIERRA ESTATES - UNIFIED INVENTORY MERGE & DEDUP")
    print("==================================================")

    records = {} # dedup_key -> unit dict
    stats = {
        'total_raw_loaded': 0,
        'duplicates_merged': 0,
        'sources': {}
    }

    # ----------------------------------------------------
    # SOURCE 1: Sierra_Estates_Owners_Units_Rent_and_Resale.xlsx (585 verified owners)
    # ----------------------------------------------------
    p1 = 'Sierra_Estates_Owners_Units_Rent_and_Resale.xlsx'
    if os.path.exists(p1):
        print(f"Loading {p1}...")
        wb = openpyxl.load_workbook(p1, read_only=True)
        count = 0
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row): continue
                stats['total_raw_loaded'] += 1
                count += 1
                
                deal = 'Rent' if 'rent' in sheetname.lower() or 'rent' in str(row[6]).lower() else 'Sale / Re-sale'
                phone = clean_phone(row[14])
                wa = f"https://wa.me/2{phone}" if phone else ''
                
                u = {
                    'unit_code': str(row[1] or '').strip(),
                    'advertiser_type': 'Owner',
                    'deal_type': deal,
                    'compound': normalize_compound(row[2]),
                    'location': str(row[3] or row[2] or 'New Cairo').strip(),
                    'zone': str(row[4] or 'New Cairo Prime').strip(),
                    'property_type': str(row[5] or 'Apartment').strip(),
                    'price_egp': clean_price(row[7]),
                    'price_display': str(row[8] or '').strip(),
                    'area_sqm': clean_num(row[9]),
                    'bedrooms': clean_num(row[10]),
                    'bathrooms': clean_num(row[11]),
                    'furnishing': str(row[12] or 'Standard').strip(),
                    'contact_name': str(row[13] or 'Owner Direct').strip(),
                    'phone': phone,
                    'whatsapp_link': wa,
                    'status': str(row[16] or 'Available').strip(),
                    'has_photos': 'Yes' if str(row[17]).lower() == 'yes' or row[18] else 'No',
                    'photo_urls': str(row[18] or '').strip(),
                    'description': str(row[19] or '').strip(),
                    'source': 'Verified Owners Direct',
                    'updated_at': str(row[21] or datetime.now().strftime('%Y-%m-%d')).strip(),
                    'duplicates_count': 1
                }
                
                key = build_dedup_key(u)
                if key in records:
                    merge_units(records[key], u)
                    stats['duplicates_merged'] += 1
                else:
                    records[key] = u
        stats['sources']['Verified_Owners_585'] = count
        print(f"  Loaded {count} rows from verified owners.")

    # ----------------------------------------------------
    # SOURCE 2: Inventory_with_Photos.xlsx (7,655 units with photos)
    # ----------------------------------------------------
    p2 = 'Inventory_with_Photos.xlsx'
    if os.path.exists(p2):
        print(f"Loading {p2}...")
        wb = openpyxl.load_workbook(p2, read_only=True)
        count = 0
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row): continue
                stats['total_raw_loaded'] += 1
                count += 1
                
                party = 'Owner' if 'owner' in sheetname.lower() else 'Broker'
                deal = 'Rent' if 'rent' in sheetname.lower() or 'rent' in str(row[6]).lower() else 'Sale / Re-sale'
                phone = clean_phone(row[14])
                wa = str(row[15] or '')
                if not wa and phone:
                    wa = f"https://wa.me/2{phone}"
                
                u = {
                    'unit_code': str(row[1] or '').strip(),
                    'advertiser_type': party,
                    'deal_type': deal,
                    'compound': normalize_compound(row[2]),
                    'location': str(row[3] or row[2] or 'New Cairo').strip(),
                    'zone': str(row[4] or 'New Cairo').strip(),
                    'property_type': str(row[5] or 'Apartment').strip(),
                    'price_egp': clean_price(row[7]),
                    'price_display': str(row[8] or '').strip(),
                    'area_sqm': clean_num(row[9]),
                    'bedrooms': clean_num(row[10]),
                    'bathrooms': clean_num(row[11]),
                    'furnishing': str(row[12] or 'Standard').strip(),
                    'contact_name': str(row[13] or '').strip(),
                    'phone': phone,
                    'whatsapp_link': wa,
                    'status': str(row[16] or 'Available').strip(),
                    'has_photos': 'Yes' if row[18] else 'No',
                    'photo_urls': str(row[18] or '').strip(),
                    'description': str(row[19] or '').strip(),
                    'source': f"Photos DB ({sheetname})",
                    'updated_at': str(row[21] or datetime.now().strftime('%Y-%m-%d')).strip(),
                    'duplicates_count': 1
                }
                
                key = build_dedup_key(u)
                if key in records:
                    merge_units(records[key], u)
                    stats['duplicates_merged'] += 1
                else:
                    records[key] = u
        stats['sources']['Inventory_with_Photos'] = count
        print(f"  Loaded {count} rows from photos inventory.")

    # ----------------------------------------------------
    # SOURCE 3: sierra-estates-master-inventory.xlsx (16,079 units)
    # ----------------------------------------------------
    p3 = 'apps/sierra-estates-realty/public/downloads/sierra-estates-master-inventory.xlsx'
    if os.path.exists(p3):
        print(f"Loading {p3}...")
        wb = openpyxl.load_workbook(p3, read_only=True)
        count = 0
        if 'All_Units' in wb.sheetnames:
            ws = wb['All_Units']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row): continue
                stats['total_raw_loaded'] += 1
                count += 1
                
                deal = normalize_deal(row[3])
                party = normalize_party(row[4], row[20])
                phone = clean_phone(row[16])
                wa = f"https://wa.me/2{phone}" if phone else ''
                
                u = {
                    'unit_code': str(row[1] or '').strip(),
                    'advertiser_type': party,
                    'deal_type': deal,
                    'compound': normalize_compound(row[5]),
                    'location': str(row[5] or 'New Cairo').strip(),
                    'zone': str(row[6] or 'New Cairo').strip(),
                    'property_type': str(row[7] or 'Apartment').strip(),
                    'price_egp': clean_price(row[13]),
                    'price_display': str(row[15] or '').strip(),
                    'area_sqm': clean_num(row[9]),
                    'bedrooms': clean_num(row[8]),
                    'bathrooms': 0,
                    'furnishing': str(row[12] or 'Standard').strip(),
                    'contact_name': str(row[17] or '').strip(),
                    'phone': phone,
                    'whatsapp_link': wa,
                    'status': str(row[2] or 'Available').strip(),
                    'has_photos': 'Yes' if str(row[21]).lower() in ('true', 'yes', '1') or row[23] else 'No',
                    'photo_urls': str(row[23] or '').strip(),
                    'description': str(row[19] or '').strip(),
                    'source': str(row[20] or 'Master Inventory Scraper').strip(),
                    'updated_at': str(row[18] or datetime.now().strftime('%Y-%m-%d')).strip(),
                    'duplicates_count': 1
                }
                
                key = build_dedup_key(u)
                if key in records:
                    merge_units(records[key], u)
                    stats['duplicates_merged'] += 1
                else:
                    records[key] = u
        stats['sources']['Giant_Master_Inventory'] = count
        print(f"  Loaded {count} rows from giant master inventory.")

    # ----------------------------------------------------
    # SOURCE 4: inventory_master_unified.csv (8,098 units)
    # ----------------------------------------------------
    p4 = 'inventory_master_unified.csv'
    if os.path.exists(p4):
        print(f"Loading {p4}...")
        count = 0
        with open(p4, mode='r', encoding='utf-8', errors='ignore') as f:
            reader = csv.DictReader(f)
            for row in reader:
                stats['total_raw_loaded'] += 1
                count += 1
                
                deal = normalize_deal(row.get('deal_type') or row.get('Operation') or '')
                party = normalize_party(row.get('owner_party') or row.get('Advertiser') or '')
                phone = clean_phone(row.get('mobile') or row.get('phone') or row.get('Phone') or '')
                wa = f"https://wa.me/2{phone}" if phone else ''
                
                u = {
                    'unit_code': str(row.get('code') or row.get('UnitCode') or '').strip(),
                    'advertiser_type': party,
                    'deal_type': deal,
                    'compound': normalize_compound(row.get('compound') or row.get('Compound') or ''),
                    'location': str(row.get('location') or row.get('Location') or 'New Cairo').strip(),
                    'zone': str(row.get('zone') or row.get('Zone') or 'New Cairo').strip(),
                    'property_type': str(row.get('property_type') or row.get('PropertyType') or 'Apartment').strip(),
                    'price_egp': clean_price(row.get('price_egp') or row.get('Price (EGP)') or row.get('price') or 0),
                    'price_display': str(row.get('price_display') or '').strip(),
                    'area_sqm': clean_num(row.get('space_m2') or row.get('Area (sqm)') or row.get('area') or 0),
                    'bedrooms': clean_num(row.get('bedrooms') or row.get('Bedrooms') or 0),
                    'bathrooms': clean_num(row.get('bathrooms') or row.get('Bathrooms') or 0),
                    'furnishing': str(row.get('furnished') or row.get('Furnishing') or 'Standard').strip(),
                    'contact_name': str(row.get('name') or row.get('Contact Name') or '').strip(),
                    'phone': phone,
                    'whatsapp_link': wa,
                    'status': str(row.get('availability') or row.get('status') or 'Available').strip(),
                    'has_photos': 'Yes' if row.get('photo_urls') or row.get('has_photos') == 'True' else 'No',
                    'photo_urls': str(row.get('photo_urls') or '').strip(),
                    'description': str(row.get('notes') or row.get('description') or '').strip(),
                    'source': 'Unified Inventory CSV',
                    'updated_at': str(row.get('timestamp') or datetime.now().strftime('%Y-%m-%d')).strip(),
                    'duplicates_count': 1
                }
                
                key = build_dedup_key(u)
                if key in records:
                    merge_units(records[key], u)
                    stats['duplicates_merged'] += 1
                else:
                    records[key] = u
        stats['sources']['Unified_Inventory_CSV'] = count
        print(f"  Loaded {count} rows from unified csv.")

    total_unique = len(records)
    print("\n--------------------------------------------------")
    print(f"TOTAL RAW ROWS INGESTED:   {stats['total_raw_loaded']:,}")
    print(f"DUPLICATES REMOVED/MERGED: {stats['duplicates_merged']:,}")
    print(f"TOTAL UNIQUE CLEAN UNITS:  {total_unique:,}")
    print("--------------------------------------------------\n")

    # Sort records: Owners first, then high price
    sorted_units = sorted(
        records.values(),
        key=lambda x: (
            0 if x['advertiser_type'] == 'Owner' else 1,
            -x['price_egp']
        )
    )

    # Assign Sequential Record IDs
    for idx, u in enumerate(sorted_units, start=1):
        u['record_id'] = f"SE-{idx:05d}"
        if not u['price_display'] and u['price_egp'] > 0:
            if u['deal_type'] == 'Rent':
                u['price_display'] = f"{int(u['price_egp']):,} EGP / Month"
            else:
                u['price_display'] = f"{int(u['price_egp']):,} EGP"

    # Split into categories
    owners_rent = [u for u in sorted_units if u['advertiser_type'] == 'Owner' and u['deal_type'] == 'Rent']
    owners_sale = [u for u in sorted_units if u['advertiser_type'] == 'Owner' and u['deal_type'] != 'Rent']
    brokers_rent = [u for u in sorted_units if u['advertiser_type'] != 'Owner' and u['deal_type'] == 'Rent']
    brokers_sale = [u for u in sorted_units if u['advertiser_type'] != 'Owner' and u['deal_type'] != 'Rent']

    print(f"Categorized Breakdown:")
    print(f"  - Owners Rent:   {len(owners_rent):,} units")
    print(f"  - Owners Sale:   {len(owners_sale):,} units")
    print(f"  - Brokers Rent:  {len(brokers_rent):,} units")
    print(f"  - Brokers Sale:  {len(brokers_sale):,} units")

    # ----------------------------------------------------
    # BUILD EXCEL WORKBOOK WITH OPENPYXL
    # ----------------------------------------------------
    wb_out = openpyxl.Workbook()
    # Remove default sheet
    wb_out.remove(wb_out.active)

    headers = [
        'Record ID',
        'Unit Code',
        'Advertiser Type',      # Owner vs Broker
        'Deal Type',             # Rent vs Sale / Re-sale
        'Compound / Project',
        'Location',
        'Zone / District',
        'Property Type',
        'Price (EGP)',
        'Price Display',
        'Area (sqm)',
        'Bedrooms',
        'Bathrooms',
        'Furnishing Status',
        'Contact Name',
        'Contact Phone',
        'Direct WhatsApp Link',
        'Listing Status',
        'Photos Available',
        'Photo URLs',
        'Description / Notes',
        'Source Channel',
        'Times Advertised (Duplicates Merged)'
    ]

    # STYLES CONFIG
    thin_side = Side(border_style='thin', color='CBD5E1') # Slate-300
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    def add_data_sheet(sheet_title, units_list, header_color):
        print(f"Writing sheet '{sheet_title}' ({len(units_list):,} units)...")
        ws = wb_out.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        
        # Add Header
        ws.append(headers)
        ws.row_dimensions[1].height = 28
        h_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type='solid')
        h_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        h_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = h_fill
            c.font = h_font
            c.alignment = h_align
            c.border = thin_border

        # Add Rows
        data_font = Font(name='Segoe UI', size=10)
        for r_idx, u in enumerate(units_list, start=2):
            ws.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            
            row_vals = [
                u['record_id'],
                u['unit_code'],
                u['advertiser_type'],
                u['deal_type'],
                u['compound'],
                u['location'],
                u['zone'],
                u['property_type'],
                u['price_egp'],
                u['price_display'],
                u['area_sqm'],
                u['bedrooms'],
                u['bathrooms'],
                u['furnishing'],
                u['contact_name'],
                u['phone'],
                u['whatsapp_link'],
                u['status'],
                u['has_photos'],
                u['photo_urls'],
                u['description'],
                u['source'],
                u['duplicates_count']
            ]
            ws.append(row_vals)
            
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if is_zebra:
                    cell.fill = zebra_fill
                    
                h_name = headers[col_idx - 1]
                if 'Price (EGP)' in h_name:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                elif h_name in ('Area (sqm)', 'Bedrooms', 'Bathrooms', 'Times Advertised (Duplicates Merged)'):
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif h_name in ('Record ID', 'Advertiser Type', 'Deal Type', 'Property Type', 'Listing Status', 'Photos Available'):
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        # AutoFilter & Freeze
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'

        # Column widths
        widths = {
            'A': 13, 'B': 18, 'C': 16, 'D': 16, 'E': 26, 'F': 22,
            'G': 20, 'H': 18, 'I': 16, 'J': 22, 'K': 12, 'L': 10,
            'M': 10, 'N': 18, 'O': 20, 'P': 18, 'Q': 28, 'R': 15,
            'S': 16, 'T': 28, 'U': 35, 'V': 22, 'W': 20
        }
        for col_l, w in widths.items():
            ws.column_dimensions[col_l].width = w

    # 1. SUMMARY SHEET
    print("Writing 'Summary' dashboard sheet...")
    ws_sum = wb_out.create_sheet(title='Summary', index=0)
    ws_sum.views.sheetView[0].showGridLines = True
    
    # Title
    ws_sum['B2'] = "SIERRA ESTATES — MASTER DEDUPLICATED INVENTORY DASHBOARD"
    ws_sum['B2'].font = Font(name='Segoe UI', size=16, bold=True, color='0F172A')
    ws_sum['B3'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Unique Listings: {total_unique:,} | Duplicates Resolved: {stats['duplicates_merged']:,}"
    ws_sum['B3'].font = Font(name='Segoe UI', size=11, color='64748B')

    # KPI Matrix
    kpis = [
        ("Total Unique Listings", f"{total_unique:,}", "1E293B"),
        ("Direct Owners Units", f"{len(owners_rent) + len(owners_sale):,}", "065F46"),
        ("Broker Network Units", f"{len(brokers_rent) + len(brokers_sale):,}", "1E3A8A"),
        ("Raw Rows Ingested", f"{stats['total_raw_loaded']:,}", "475569"),
        ("Duplicates Cleaned", f"{stats['duplicates_merged']:,}", "B45309"),
    ]
    
    for i, (label, val, col) in enumerate(kpis):
        c_letter = chr(ord('B') + (i * 2))
        ws_sum[f'{c_letter}5'] = label
        ws_sum[f'{c_letter}5'].font = Font(name='Segoe UI', size=10, bold=True, color='64748B')
        ws_sum[f'{c_letter}6'] = val
        ws_sum[f'{c_letter}6'].font = Font(name='Segoe UI', size=20, bold=True, color=col)

    # Categories Table
    ws_sum['B9'] = "Inventory Portfolio Breakdown by Sheet"
    ws_sum['B9'].font = Font(name='Segoe UI', size=13, bold=True, color='0F172A')
    
    cat_headers = ["Sheet Name", "Classification", "Deal Type", "Total Units", "% of Portfolio", "Average Price"]
    for col_i, h in enumerate(cat_headers, start=2):
        cell = ws_sum.cell(row=11, column=col_i, value=h)
        cell.font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    cat_data = [
        ("Owners_Rent", "Direct Owner", "Rent", len(owners_rent), f"{len(owners_rent)/total_unique*100:.1f}%", f"{sum(u['price_egp'] for u in owners_rent)/(len(owners_rent) or 1):,.0f} EGP"),
        ("Owners_Sale", "Direct Owner", "Sale / Re-sale", len(owners_sale), f"{len(owners_sale)/total_unique*100:.1f}%", f"{sum(u['price_egp'] for u in owners_sale)/(len(owners_sale) or 1):,.0f} EGP"),
        ("Brokers_Rent", "Broker Network", "Rent", len(brokers_rent), f"{len(brokers_rent)/total_unique*100:.1f}%", f"{sum(u['price_egp'] for u in brokers_rent)/(len(brokers_rent) or 1):,.0f} EGP"),
        ("Brokers_Sale", "Broker Network", "Sale / Re-sale", len(brokers_sale), f"{len(brokers_sale)/total_unique*100:.1f}%", f"{sum(u['price_egp'] for u in brokers_sale)/(len(brokers_sale) or 1):,.0f} EGP"),
        ("All_Units", "Consolidated Agency Master", "All Types", total_unique, "100.0%", "—"),
    ]

    for row_offset, row_data in enumerate(cat_data, start=12):
        for col_i, val in enumerate(row_data, start=2):
            cell = ws_sum.cell(row=row_offset, column=col_i, value=val)
            cell.font = Font(name='Segoe UI', size=10)
            cell.border = thin_border
            if row_offset % 2 == 1:
                cell.fill = zebra_fill
            cell.alignment = Alignment(horizontal='center' if col_i in (3, 4, 5, 6) else 'left')

    # Top Compounds Table
    compound_counts = {}
    for u in sorted_units:
        c_name = u['compound']
        compound_counts[c_name] = compound_counts.get(c_name, 0) + 1
    top_compounds = sorted(compound_counts.items(), key=lambda x: -x[1])[:12]

    ws_sum['B19'] = "Top 12 Compound Communities by Units"
    ws_sum['B19'].font = Font(name='Segoe UI', size=13, bold=True, color='0F172A')

    ws_sum['B21'] = "Rank"
    ws_sum['C21'] = "Compound Name"
    ws_sum['D21'] = "Unique Units"
    ws_sum['E21'] = "Market Share"
    for col_i in range(2, 6):
        cell = ws_sum.cell(row=21, column=col_i)
        cell.font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid') # Teal 700
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for rank, (comp_name, count) in enumerate(top_compounds, start=1):
        r = 21 + rank
        ws_sum.cell(row=r, column=2, value=rank).alignment = Alignment(horizontal='center')
        ws_sum.cell(row=r, column=3, value=comp_name).alignment = Alignment(horizontal='left')
        ws_sum.cell(row=r, column=4, value=f"{count:,}").alignment = Alignment(horizontal='center')
        ws_sum.cell(row=r, column=5, value=f"{count/total_unique*100:.1f}%").alignment = Alignment(horizontal='center')
        for c in range(2, 6):
            cell = ws_sum.cell(row=r, column=c)
            cell.font = Font(name='Segoe UI', size=10)
            cell.border = thin_border
            if rank % 2 == 0:
                cell.fill = zebra_fill

    ws_sum.column_dimensions['A'].width = 4
    ws_sum.column_dimensions['B'].width = 24
    ws_sum.column_dimensions['C'].width = 30
    ws_sum.column_dimensions['D'].width = 18
    ws_sum.column_dimensions['E'].width = 18
    ws_sum.column_dimensions['F'].width = 20
    ws_sum.column_dimensions['G'].width = 18

    # 2. ALL_UNITS SHEET
    add_data_sheet('All_Units', sorted_units, '1E293B') # Slate Navy

    # 3. OWNERS_RENT SHEET
    add_data_sheet('Owners_Rent', owners_rent, '065F46') # Emerald Green

    # 4. OWNERS_SALE SHEET
    add_data_sheet('Owners_Sale', owners_sale, '78350F') # Warm Gold / Amber

    # 5. BROKERS_RENT SHEET
    add_data_sheet('Brokers_Rent', brokers_rent, '115E59') # Deep Teal

    # 6. BROKERS_SALE SHEET
    add_data_sheet('Brokers_Sale', brokers_sale, '312E81') # Royal Indigo

    # SAVE TO CANONICAL LOCATIONS
    out_xlsx = 'Sierra_Estates_Master_Database_Unified.xlsx'
    print(f"\nSaving final styled workbook to {out_xlsx}...")
    wb_out.save(out_xlsx)
    print(f"Saved: {out_xlsx}")

    # Copy to public/downloads for live web access
    pub_path = os.path.join('apps', 'sierra-estates-realty', 'public', 'downloads', out_xlsx)
    import shutil
    shutil.copy(out_xlsx, pub_path)
    print(f"Copied to: {pub_path}")

    # Also export individual CSVs for each tab
    print("\nExporting individual CSVs for each sheet...")
    for sheet in wb_out.sheetnames:
        if sheet == 'Summary': continue
        csv_name = f"Sierra_Estates_{sheet}.csv"
        ws_curr = wb_out[sheet]
        with open(csv_name, 'w', newline='', encoding='utf-8-sig') as f_csv:
            writer = csv.writer(f_csv)
            for row in ws_curr.iter_rows(values_only=True):
                writer.writerow(row)
        shutil.copy(csv_name, os.path.join('apps', 'sierra-estates-realty', 'public', 'downloads', csv_name))
        print(f"  Exported CSV: {csv_name}")

    print("\n==================================================")
    print("SUCCESS: ALL INVENTORY CONSOLIDATED WITHOUT DUPLICATION")
    print("==================================================")

if __name__ == '__main__':
    main()
