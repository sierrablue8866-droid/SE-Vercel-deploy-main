from __future__ import annotations

import csv
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path('/home/ubuntu')
PROFILE = ROOT / 'whatsapp_inventory/new_owners_aug2026/profile'
XLSX = ROOT / 'whatsapp_inventory/merged_inventory/Sierra_Estates_All_Inventory_One_Sheet.xlsx'
CSV_OUT = ROOT / 'whatsapp_inventory/merged_inventory/Sierra_Estates_All_Inventory_Airtable.csv'
rows = json.loads((PROFILE / 'airtable_new_records.json').read_text(encoding='utf-8'))

# Append to the single-sheet workbook, preserving its existing formatting/table.
wb = load_workbook(XLSX)
ws = wb['All Inventory Units']
header_row = 5
headers = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
existing_ids = {str(ws.cell(row, 1).value).strip() for row in range(header_row + 1, ws.max_row + 1) if ws.cell(row, 1).value}
new_rows = [row for row in rows if row['RecordID'] not in existing_ids]
for row_index, row in enumerate(new_rows, header_row + 1 + len(existing_ids) - len(existing_ids)):
    # Use the actual next worksheet row rather than relying on the ID count.
    row_index = ws.max_row + 1
    for col, field in enumerate(headers, 1):
        value = row.get(field, '')
        ws.cell(row_index, col, value)
        ws.cell(row_index, col).alignment = ws.cell(header_row + 1, col).alignment.copy()
        ws.cell(row_index, col).font = ws.cell(header_row + 1, col).font.copy()
    for col in range(1, len(headers) + 1):
        if row_index % 2 == 0:
            ws.cell(row_index, col).fill = ws.cell(header_row + 2, col).fill.copy()
# Extend table range if present.
if ws.tables:
    table = next(iter(ws.tables.values()))
    table.ref = f'A{header_row}:{ws.cell(header_row, ws.max_column).column_letter}{ws.max_row}'
wb.save(XLSX)

# Append the same canonical fields to the Airtable CSV export.
with CSV_OUT.open(encoding='utf-8-sig', newline='') as handle:
    reader = csv.DictReader(handle)
    fieldnames = reader.fieldnames or []
    existing_csv = list(reader)
existing_csv_ids = {str(r.get('RecordID') or '').strip() for r in existing_csv}
for row in new_rows:
    payload = {field: row.get(field, '') for field in fieldnames}
    if row['RecordID'] not in existing_csv_ids:
        existing_csv.append(payload)
with CSV_OUT.open('w', encoding='utf-8-sig', newline='') as handle:
    writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(existing_csv)

print(f'new_records_appended={len(new_rows)}')
print(f'workbook_rows={ws.max_row - header_row}')
print(f'csv_rows={len(existing_csv)}')
