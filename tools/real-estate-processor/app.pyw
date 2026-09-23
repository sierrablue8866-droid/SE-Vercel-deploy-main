# -*- coding: utf-8 -*-
"""Real Estate Processor - one file, GUI, no terminal.
FIXED: advertiser detection now checks folder path too ('Active owners' folder -> Owner),
'chat' no longer forces Broker, and output splits into:
Owners_Rent / Owners_Sale / Brokers_Rent / Brokers_Sale / Undefined.
"""
import os, sys, re, glob, time, subprocess, threading, importlib
import tkinter as tk
from tkinter import filedialog, messagebox

pd = None
np = None
try:
    import pandas as pd
    import numpy as np
except ImportError:
    pass

FROZEN = bool(getattr(sys, "frozen", False))
USD_TO_EGP = 48.0
ID_PREFIX = "SB"
SKIP_FILES = ("Final_", "~$", "Missing_")
SKIP_SHEETS = ("dashboard", "summary", "pivot", "تعليمات")
UNK = {"*", "", "nan", "none", "null", "-", "n/a", "unknown", "غير مذكور"}

SYN = {
    "Unit_Code": ["code", "كود", "unit code", "رقم الوحدة", "ref"],
    "Phone": ["mobile", "phone", "تليفون", "موبايل", "رقم الهاتف", "broker phone", "contact", "whatsapp"],
    "Price_Raw": ["price", "السعر", "الإيجار", "ايجار", "المطلوب", "unit price", "total price"],
    "Owner_Name": ["owner name", "اسم المالك", "الاسم", "name"],
    "Contact_Name": ["broker name", "اسم البروكر", "المعلن", "contact person"],
    "Listing_Date": ["timestamp", "listing date", "تاريخ العرض", "تاريخ الإعلان", "date"],
    "Update_Date": ["تاريخ اخر تحديث", "update date", "last update"],
    "Availability": ["availability", "avail", "الحالة", "المتاحية", "status"],
    "Rooms": ["bedrooms", "عدد الغرف", "الغرف", "rooms", "نوم"],
    "Bathrooms": ["bathrooms", "baths", "الحمامات", "حمام"],
    "Location": ["location", "المنطقة والكمبوند", "الكمبوند", "الموقع", "compound"],
    "Furnishing": ["furnishing", "حالة التأثيث", "التأثيث", "مفروش"],
    "Finishing": ["finishing", "تشطيب", "حالة التشطيب"],
    "Unit_Type": ["property type", "unit type", "نوع الوحدة", "النوع"],
    "Deal": ["transaction", "نوع المعاملة", "deal", "بيع/ايجار"],
    "Advertiser_Type": ["advertiser type", "نوع المعلن", "owner/broker", "المعلن"],
    "Area": ["space", "area", "المساحة", "مساحه"],
    "Source_Group": ["source_sheet", "group name", "اسم الجروب", "المصدر"],
    "Notes": ["notes", "ملاحظات", "تفاصيل", "الوصف", "description"],
}

COMPOUNDS = {
    "Madinaty": [r"مدينت[يى]", r"madinat"], "Al Rehab": [r"الرحاب", r"rehab"],
    "Mivida": [r"ميفيدا", r"mivida"], "Hyde Park": [r"هايد\s*بارك", r"hyde\s*park"],
    "Mountain View": [r"ماونتن\s*فيو", r"mountain\s*view"], "Villette": [r"فيليت", r"villette"],
    "Palm Hills": [r"بالم\s*هيلز", r"palm\s*hills"], "Eastown": [r"ايست\s*تاون", r"eastown"],
    "Cairo Festival": [r"فستيفال", r"\bcfc\b"], "Fifth Square": [r"فيفت\s*سكوير", r"fifth\s*square"],
    "Sodic": [r"سوديك", r"sodic"], "New Cairo": [r"التجمع", r"new\s*cairo"],
    "Sheikh Zayed": [r"الشيخ\s*زايد", r"zayed"], "North Coast": [r"الساحل\s*الشمالي", r"north\s*coast"],
    "El Shorouk": [r"الشروق", r"shorouk"],
}

# Advertiser keywords (priority order matters - see detection below)
RE_OWNER = r"مالك|ملاك|اونر|أونر|owner|من\s*المالك"
RE_BROKER = r"بروكر|وسيط|سمسار|شرك[ةه]|تسويق|مكتب|broker|agent|agency|realtor"

# ---------------- helpers ----------------
def clean_val(v):
    if pd.isna(v):
        return "Unknown"
    s = str(v).strip()
    return "Unknown" if s.lower() in UNK else s

def norm_phone(v):
    if pd.isna(v):
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    d = re.sub(r"\D", "", s)
    if not d:
        return None
    if d.startswith("20") and len(d) >= 12:
        d = d[2:]
    if len(d) == 10 and d.startswith("1"):
        d = "0" + d
    return d if len(d) >= 7 else None

def last7(v):
    d = re.sub(r"\D", "", str(v)) if v else ""
    return d[-7:] if len(d) >= 7 else ""

def parse_price(v):
    if pd.isna(v):
        return (np.nan, "EGP")
    s = str(v).strip()
    if s.lower() in UNK:
        return (np.nan, "EGP")
    cur = "USD" if re.search(r"\$|usd|dollar|دولار", s, re.I) else "EGP"
    m = re.search(r"\d+(?:[.,]\d+)*", s)
    if not m:
        return (np.nan, cur)
    t = m.group(0)
    if re.fullmatch(r"\d{1,3}(?:[.,]\d{3})+", t):
        t = re.sub(r"[.,]", "", t)
    else:
        t = t.replace(",", "")
    try:
        n = float(t)
    except ValueError:
        return (np.nan, cur)
    tail = s[m.end():m.end() + 10]
    if re.match(r"\s*(مليون|ملون|million)\b", tail, re.I) or re.match(r"\s*m(?![2²0-9a-z])", tail, re.I):
        n *= 1_000_000
    elif re.match(r"\s*(الف|ألف|thousand)\b", tail, re.I) or re.match(r"\s*k(?![a-z0-9])", tail, re.I):
        n *= 1_000
    return (n, cur)

def parse_num(v):
    if pd.isna(v):
        return np.nan
    s = str(v).strip().lower()
    if s in UNK:
        return np.nan
    m = re.search(r"\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else np.nan

def parse_date(v):
    if pd.isna(v):
        return pd.NaT
    s = str(v).strip()
    if s.lower() in ("", "nan", "none", "nat", "*", "unknown"):
        return pd.NaT
    for kw in ({"dayfirst": True}, {"dayfirst": False}, {"format": "mixed"}):
        try:
            return pd.to_datetime(s, errors="raise", **kw)
        except Exception:
            pass
    return pd.NaT

def norm_code(v):
    s = clean_val(v)
    if s == "Unknown":
        return np.nan
    s = re.sub(r"^(كود|code|unit|ref)\s*[:\-_]?\s*", "", s, flags=re.I)
    s = re.sub(r"\s+", "", s).upper()
    return s if s else np.nan

def norm_deal(v):
    s = str(v).strip().lower() if pd.notna(v) else ""
    if any(w in s for w in ("ايجار", "إيجار", "rent")):
        return "Rent"
    if any(w in s for w in ("بيع", "sale", "sell", "resale", "تنازل")):
        return "Sale"
    return "Unknown"

def norm_compound(loc, notes):
    comb = (str(loc) + " " + str(notes)).lower()
    for std, pats in COMPOUNDS.items():
        if any(re.search(p, comb, re.I) for p in pats):
            return std
    return clean_val(loc)

def map_cols(headers):
    lows = [str(h).strip().lower() if pd.notna(h) else "" for h in headers]
    m, used = {}, set()
    for std, syns in SYN.items():
        for syn in syns:
            for i, h in enumerate(lows):
                if i in used or not h:
                    continue
                if syn in h:
                    m[std] = i
                    used.add(i)
                    break
        if std in m:
            break
    return m

def find_header(raw, scan=10):
    for i in range(min(scan, len(raw))):
        m = map_cols(raw.iloc[i].tolist())
        if len(m) >= 2 and any(k in m for k in ("Phone", "Unit_Code", "Price_Raw")):
            return i, m
    return None, None

# ---------------- WhatsApp ----------------
WA_MSG = re.compile(r"^\[?(\d{1,2}[\/.\-]\d{1,2}[\/.\-]\d{2,4}),?\s+(\d{1,2}:\d{2}(?::\d{2})?\s*(?:[AaPp][Mm]|ص|م)?)[\]\s\-:]+([^:]+):")
WA_PHONE = re.compile(r"(?<!\d)(?:\+?20[\s\-.]?)?0?1[0125](?:[\s\-.]?\d){8}(?!\d)")

def wa_price(text):
    m = re.search(r"(\d+(?:[.,]\d+)*)\s*(?:مليون|ملون|million|الف|ألف|k\b|جنيه|egp|usd|\$|دولار)", text, re.I)
    if m:
        return parse_price(m.group(0))
    for n in re.findall(r"\b\d{3,8}\b", text):
        if not 1990 <= int(n) <= 2035:
            return parse_price(n)
    return (np.nan, "EGP")

def wa_msg_rec(text, dt, sender, group, fp, fmod):
    text = text.strip()
    if len(text) < 15:
        return None
    pm = WA_PHONE.search(text)
    phone = norm_phone(pm.group(0)) if pm else norm_phone(re.sub(r"\D", "", sender))
    if not phone:
        return None
    pv, cur = wa_price(text)
    cm = re.search(r"(?:كود|code|ref)\s*[:\-_]?\s*([A-Za-z0-9\-_]+)", text, re.I)
    am = re.search(r"(\d{2,4})\s*(?:متر|م²|m2|sqm)", text, re.I)
    rm = re.search(r"(\d)\s*(?:غرف|غرفة|نوم|rooms?|beds?)", text, re.I)
    return {
        "Unit_Code": norm_code(cm.group(1)) if cm else np.nan,
        "Phone": phone, "Price_Raw": pv, "Currency": cur,
        "Price_In_Thousands": False,
        "Owner_Name": sender, "Contact_Name": sender,
        "Listing_Date": dt, "Update_Date": dt, "Availability": "Available",
        "Rooms": float(rm.group(1)) if rm else np.nan, "Bathrooms": np.nan,
        "Area": float(am.group(1)) if am else np.nan,
        "Location": norm_compound("", text),
        "Furnishing": "Furnished" if "مفروش" in text else ("Not Furnished" if "غير مفروش" in text else "Unknown"),
        "Finishing": "Unknown", "Unit_Type": "Unknown",
        "Deal": norm_deal(text), "Advertiser_Type": "Unknown",
        "Source_Group": group, "Notes": text[:500],
        "Source_File": os.path.basename(fp), "Source_Sheet": "WhatsApp",
        "File_Modified": fmod, "Full_Path": fp,
    }

def parse_whatsapp(src, log):
    recs = []
    for fp in glob.glob(os.path.join(src, "**", "*.txt"), recursive=True):
        fn = os.path.basename(fp)
        if any(fn.startswith(p) for p in SKIP_FILES):
            continue
        group = os.path.splitext(fn)[0]
        fmod = pd.to_datetime(os.path.getmtime(fp), unit="s")
        with open(fp, "r", encoding="utf-8", errors="ignore") as f:
            cm, cd, cs = "", fmod, "Unknown"
            for line in f:
                m = WA_MSG.match(line)
                if m:
                    if cm:
                        r = wa_msg_rec(cm, cd, cs, group, fp, fmod)
                        if r:
                            recs.append(r)
                    d, t, cs = m.groups()
                    cd = parse_date(d + " " + t) or fmod
                    cm = line[m.end():].strip()
                else:
                    cm += " " + line.strip()
            if cm:
                r = wa_msg_rec(cm, cd, cs, group, fp, fmod)
                if r:
                    recs.append(r)
    if recs:
        log("  WhatsApp: %d messages\n" % len(recs))
    return recs

# ---------------- Pipeline ----------------
def run_pipeline(src, out_dir, log):
    global pd, np
    if pd is None or np is None:
        import pandas, numpy
        pd, np = pandas, numpy
    t0 = time.time()

    dfs = []
    for p in glob.glob(os.path.join(src, "**", "*.xls*"), recursive=True):
        fn = os.path.basename(p)
        if any(fn.startswith(s) for s in SKIP_FILES):
            continue
        try:
            sheets = pd.read_excel(p, sheet_name=None, header=None)
        except Exception as e:
            log("  skip %s: %s\n" % (fn, e))
            continue
        for sn, raw in sheets.items():
            if any(s in str(sn).lower() for s in SKIP_SHEETS):
                continue
            idx, mp = find_header(raw)
            if idx is None:
                continue
            data = raw.iloc[idx + 1:]
            if data.empty:
                continue
            out = pd.DataFrame(index=range(len(data)))
            for std, pos in mp.items():
                out[std] = data.iloc[:, pos].values
            ph = str(raw.iloc[idx, mp["Price_Raw"]]).lower() if "Price_Raw" in mp else ""
            out["Price_In_Thousands"] = any(w in ph for w in ("thousand", "'000", "بالألف"))
            out["Source_File"] = fn
            out["Source_Sheet"] = str(sn)
            out["File_Modified"] = pd.to_datetime(os.path.getmtime(p), unit="s")
            out["Full_Path"] = p          # <- needed for folder-based advertiser detection
            dfs.append(out)
            log("  %s | %s: %d rows\n" % (fn, sn, len(out)))

    wa = parse_whatsapp(src, log)
    if wa:
        dfs.append(pd.DataFrame(wa))
    if not dfs:
        return {"status": "error", "message": "No data found in that folder."}

    df = pd.concat(dfs, ignore_index=True, sort=False)
    total = len(df)
    log("\nLoaded %d rows\n" % total)

    for c in SYN:
        if c not in df.columns:
            df[c] = np.nan
    for c in ("Currency", "Full_Path"):
        if c not in df.columns:
            df[c] = ""
    if "Price_In_Thousands" not in df.columns:
        df["Price_In_Thousands"] = False
    df["Price_In_Thousands"] = df["Price_In_Thousands"].fillna(False)

    df["Phone_Clean"] = df["Phone"].map(norm_phone)
    df = df[df["Phone_Clean"].notna()].copy()
    dropped = total - len(df)
    df["Phone_Last7"] = df["Phone_Clean"].map(last7)
    log("With phone: %d (dropped %d)\n" % (len(df), dropped))

    usd_hint = df["Currency"].astype(str).str.upper().eq("USD")
    pr = df["Price_Raw"].map(parse_price)
    df["Price"] = [t[0] for t in pr]
    df["Currency"] = [t[1] for t in pr]
    df["Currency"] = np.where(usd_hint, "USD", df["Currency"])
    k = df["Price_In_Thousands"] & df["Price"].notna()
    df.loc[k, "Price"] = df.loc[k, "Price"] * 1000

    deal_raw = df["Deal"].fillna("").astype(str).str.lower()
    is_sale = deal_raw.str.contains("بيع|sale|sell|resale|تنازل", regex=True)
    is_rent = deal_raw.str.contains("ايجار|إيجار|rent", regex=True)
    egp = df["Currency"].eq("EGP") & ~df["Price_In_Thousands"]
    m = egp & is_sale & df["Price"].between(0.01, 99.99)
    df.loc[m, "Price"] = df.loc[m, "Price"] * 1_000_000
    m = egp & is_rent & df["Price"].between(0.01, 99.99)
    df.loc[m, "Price"] = df.loc[m, "Price"] * 1_000
    df["Price_EGP"] = np.where(df["Currency"].eq("USD"), df["Price"] * USD_TO_EGP, df["Price"])
    df["Deal_Clean"] = df["Deal"].map(norm_deal)

    # Deal fallback by price (so fewer rows land in Undefined)
    unk = df["Deal_Clean"].eq("Unknown") & df["Price_EGP"].notna()
    df.loc[unk & (df["Price_EGP"] > 500_000), "Deal_Clean"] = "Sale"
    df.loc[unk & df["Price_EGP"].between(0.01, 500_000), "Deal_Clean"] = "Rent"

    # ==================== ADVERTISER DETECTION (FIXED) ====================
    # Priority: 1) explicit column  2) filename+sheet+group  3) FULL PATH (folder!)
    #           4) notes.  'chat' no longer forces Broker.
    val    = df["Advertiser_Type"].fillna("").astype(str).str.lower()
    spec   = (df["Source_File"].fillna("").astype(str) + " " +
              df["Source_Sheet"].fillna("").astype(str) + " " +
              df["Source_Group"].fillna("").astype(str)).str.lower()
    folder = df["Full_Path"].fillna("").astype(str).str.lower()
    notes_l = df["Notes"].fillna("").astype(str).str.lower()

    df["Adv_Clean"] = np.select(
        [val.str.contains(RE_OWNER, regex=True),
         val.str.contains(RE_BROKER, regex=True),
         spec.str.contains(RE_OWNER, regex=True),
         spec.str.contains(RE_BROKER, regex=True),
         folder.str.contains(RE_OWNER, regex=True),
         folder.str.contains(RE_BROKER, regex=True),
         notes_l.str.contains(RE_OWNER, regex=True),
         notes_l.str.contains(RE_BROKER, regex=True)],
        ["Owner", "Broker", "Owner", "Broker", "Owner", "Broker", "Owner", "Broker"],
        default="Unknown")
    log("Advertisers -> Owners: %d | Brokers: %d | Undefined: %d\n" % (
        (df["Adv_Clean"] == "Owner").sum(),
        (df["Adv_Clean"] == "Broker").sum(),
        (df["Adv_Clean"] == "Unknown").sum()))

    notes = df["Notes"].fillna("").astype(str)
    df["Loc_Clean"] = [norm_compound(a, b) for a, b in zip(df["Location"].tolist(), notes.tolist())]
    df["Code_Clean"] = df["Unit_Code"].map(norm_code)
    df["Rooms_N"] = df["Rooms"].map(parse_num)
    df["Baths_N"] = df["Bathrooms"].map(parse_num)
    df["Area_N"] = df["Area"].map(parse_num)
    df["Furn_Clean"] = df["Furnishing"].map(clean_val)
    df["Fin_Clean"] = df["Finishing"].map(clean_val)
    df["Type_Clean"] = df["Unit_Type"].map(clean_val)
    df["Owner_Clean"] = df["Owner_Name"].map(clean_val)
    df["Contact_Clean"] = df["Contact_Name"].map(clean_val)
    df["Notes_Clean"] = df["Notes"].map(clean_val)
    df["Avail_Clean"] = df["Availability"].map(clean_val)
    df["LDate"] = df["Listing_Date"].map(parse_date).fillna(df["File_Modified"])
    df["UDate"] = df["Update_Date"].map(parse_date).fillna(df["LDate"])

    comb = (deal_raw + " " + notes).str.lower()
    df["Status"] = np.where(comb.str.contains("sold|تم البيع|اتباعت", regex=True), "Sold",
                   np.where(comb.str.contains("rented|تم الايجار|تم التأجير", regex=True), "Rented", "Active"))

    df["Key"] = (df["Phone_Last7"] + "|" +
                 df["Price_EGP"].round(0).astype("Int64").astype(str) + "|" +
                 df["Deal_Clean"])
    df = df.sort_values(["UDate", "LDate"], ascending=False)
    st = df.groupby("Key").agg(First=("LDate", "min"), Last=("UDate", "max"), N=("Key", "size")).reset_index()
    codes = (df.dropna(subset=["Code_Clean"]).groupby("Key")["Code_Clean"]
             .agg(lambda x: ", ".join(sorted(set(str(i) for i in x))))
             .rename("All_Codes").reset_index())
    u = df.drop_duplicates("Key", keep="first").merge(st, on="Key").merge(codes, on="Key", how="left")
    u["All_Codes"] = u["All_Codes"].fillna(u["Code_Clean"])
    u["Days"] = (u["Last"] - u["First"]).dt.days
    u = u.sort_values("Last", ascending=False).reset_index(drop=True)
    u["Unit_ID"] = ["%s-%04d" % (ID_PREFIX, i + 1) for i in range(len(u))]
    dups = len(df) - len(u)
    log("Unique units: %d (merged %d duplicates)\n" % (len(u), dups))

    if not os.path.isdir(out_dir):
        out_dir = os.path.expanduser("~")
    path = os.path.join(out_dir, "Final_RealEstate_Database.xlsx")
    COLS = {
        "Unit_ID": "Unit_ID", "Phone_Clean": "Phone", "Phone_Last7": "Phone_Last7",
        "Deal_Clean": "Deal", "Status": "Status", "Adv_Clean": "Advertiser_Type",
        "Avail_Clean": "Availability", "Price_EGP": "Price_EGP", "Price": "Price_Raw",
        "Currency": "Currency", "Loc_Clean": "Compound", "Type_Clean": "Unit_Type",
        "Area_N": "Area_m2", "Rooms_N": "Rooms", "Baths_N": "Bathrooms",
        "Furn_Clean": "Furnishing", "Fin_Clean": "Finishing",
        "Owner_Clean": "Owner_Name", "Contact_Clean": "Contact_Person",
        "Code_Clean": "Unit_Code", "All_Codes": "All_Original_Codes",
        "First": "First_Seen", "Last": "Latest_Update", "Days": "Days_Advertised",
        "N": "Listings_Count", "Source_File": "Source_File", "Source_Sheet": "Source_Sheet",
        "Notes_Clean": "Notes",
    }
    master = u[[c for c in COLS if c in u.columns]].rename(columns=COLS)

    def write_book(p):
        with pd.ExcelWriter(p, engine="openpyxl") as w:
            pd.DataFrame({
                "Metric": ["Total rows", "Dropped (no phone)", "Unique units", "Duplicates merged",
                           "Owners Rent", "Owners Sale", "Brokers Rent", "Brokers Sale", "Undefined",
                           "Active", "Sold", "Rented", "Seconds"],
                "Value": [total, dropped, len(master), dups,
                          len(master[(master["Advertiser_Type"] == "Owner") & (master["Deal"] == "Rent")]),
                          len(master[(master["Advertiser_Type"] == "Owner") & (master["Deal"] == "Sale")]),
                          len(master[(master["Advertiser_Type"] == "Broker") & (master["Deal"] == "Rent")]),
                          len(master[(master["Advertiser_Type"] == "Broker") & (master["Deal"] == "Sale")]),
                          len(master[~(((master["Advertiser_Type"] == "Owner") | (master["Advertiser_Type"] == "Broker"))
                                       & ((master["Deal"] == "Rent") | (master["Deal"] == "Sale")))]),
                          (master["Status"] == "Active").sum(),
                          (master["Status"] == "Sold").sum(),
                          (master["Status"] == "Rented").sum(),
                          round(time.time() - t0, 1)],
            }).to_excel(w, sheet_name="Summary", index=False)
            master.to_excel(w, sheet_name="All_Units", index=False)

            # ---- 5 split sheets (always written, even if empty) ----
            masks = {
                "Owners_Rent":   (master["Advertiser_Type"] == "Owner") & (master["Deal"] == "Rent"),
                "Owners_Sale":   (master["Advertiser_Type"] == "Owner") & (master["Deal"] == "Sale"),
                "Brokers_Rent":  (master["Advertiser_Type"] == "Broker") & (master["Deal"] == "Rent"),
                "Brokers_Sale":  (master["Advertiser_Type"] == "Broker") & (master["Deal"] == "Sale"),
            }
            defined = pd.Series(False, index=master.index)
            for msk in masks.values():
                defined = defined | msk
            masks["Undefined"] = ~defined   # anything we can't classify (owner/broker OR deal unknown)

            for name, msk in masks.items():
                master[msk].to_excel(w, sheet_name=name, index=False)
        style_file(p)

    try:
        write_book(path)
    except PermissionError:
        path = os.path.join(out_dir, "Final_RealEstate_Database_%d.xlsx" % int(time.time()))
        write_book(path)

    dur = round(time.time() - t0, 1)
    log("\nDone in %ss\nFile: %s\n" % (dur, path))
    return {"status": "success", "output_file": path, "units": len(master), "dups": dups}

def style_file(path):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = openpyxl.load_workbook(path)
        hf = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        hfont = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        dfont = Font(name="Segoe UI", size=9)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(ws.max_column), ws.max_row)
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(1, c)
                cell.fill = hf
                cell.font = hfont
                cell.alignment = Alignment(horizontal="center", vertical="center")
                head = str(cell.value or "").lower()
                for r in range(2, ws.max_row + 1):
                    cl = ws.cell(r, c)
                    cl.font = dfont
                    if "price" in head and isinstance(cl.value, (int, float)):
                        cl.number_format = "#,##0"
                mx = max((len(str(ws.cell(r, c).value or "")) for r in range(1, min(ws.max_row, 60) + 1)), default=8)
                ws.column_dimensions[get_column_letter(c)].width = min(mx + 3, 42)
        wb.save(path)
    except Exception as e:
        print("style skipped:", e)

# ---------------- GUI ----------------
def check_deps():
    miss = []
    for mod in ("numpy", "pandas", "openpyxl"):
        try:
            importlib.import_module(mod)
        except ImportError:
            miss.append(mod)
    return miss

def install_deps(log):
    if FROZEN:
        log("Bundled EXE - nothing to install.\n")
        return True
    miss = check_deps()
    if not miss:
        log("All dependencies ready.\n")
        return True
    log("Installing %s ... (1-2 min)\n" % ", ".join(miss))
    try:
        kw = {"creationflags": 0x08000000} if os.name == "nt" else {}
        subprocess.check_call([sys.executable, "-m", "pip", "install"] + miss, **kw)
        log("Installed.\n")
        return True
    except Exception as e:
        log("Install failed: %s\n" % e)
        return False

class App:
    def __init__(self, root):
        self.root = root
        self.result = None
        root.title("Real Estate Processor")
        root.geometry("760x620")
        root.configure(bg="#f0f2f5")
        tk.Frame(root, bg="#1F497D", height=52).pack(fill="x")
        tk.Label(root, text="🏠  Real Estate Processor", fg="white", bg="#1F497D",
                 font=("Segoe UI", 16, "bold")).pack(pady=10)
        f = tk.Frame(root, bg="#f0f2f5", padx=20, pady=10)
        f.pack(fill="both", expand=True)

        tk.Label(f, text="Source folder:", bg="#f0f2f5", font=("Segoe UI", 10, "bold")).grid(row=0, column=0, sticky="w", pady=6)
        self.src = tk.StringVar(value=r"H:\Sheets\Active owners")
        tk.Entry(f, textvariable=self.src, width=54).grid(row=0, column=1, pady=6, padx=8)
        tk.Button(f, text="📁", command=lambda: self.pick(self.src)).grid(row=0, column=2)

        tk.Label(f, text="Output folder:", bg="#f0f2f5", font=("Segoe UI", 10, "bold")).grid(row=1, column=0, sticky="w", pady=6)
        self.out = tk.StringVar(value="H:\\" if os.path.isdir("H:\\") else os.path.expanduser("~"))
        tk.Entry(f, textvariable=self.out, width=54).grid(row=1, column=1, pady=6, padx=8)
        tk.Button(f, text="📁", command=lambda: self.pick(self.out)).grid(row=1, column=2)

        b = tk.Frame(f, bg="#f0f2f5")
        b.grid(row=2, column=0, columnspan=3, pady=12)
        self.run_btn = tk.Button(b, text="▶  Run", command=self.run, font=("Segoe UI", 13, "bold"),
                                 bg="#16A34A", fg="white", padx=28, pady=8, relief="flat", cursor="hand2")
        self.run_btn.pack(side="left", padx=6)
        if not FROZEN:
            tk.Button(b, text="📦 Install", command=self.deps, font=("Segoe UI", 10),
                      bg="#2563EB", fg="white", padx=14, pady=8, relief="flat", cursor="hand2").pack(side="left", padx=6)
        tk.Button(b, text="📂 Open Result", command=self.open_result, font=("Segoe UI", 10),
                  bg="#6B7280", fg="white", padx=14, pady=8, relief="flat", cursor="hand2").pack(side="left", padx=6)

        tk.Label(f, text="Log:", bg="#f0f2f5", font=("Segoe UI", 10, "bold")).grid(row=3, column=0, sticky="w", pady=(6, 0))
        lf = tk.Frame(f)
        lf.grid(row=4, column=0, columnspan=3, sticky="nsew", pady=4)
        f.rowconfigure(4, weight=1)
        f.columnconfigure(1, weight=1)
        self.log = tk.Text(lf, height=14, font=("Consolas", 9), bg="#1e1e1e", fg="#4ade80",
                           insertbackground="white", wrap="word", relief="flat", padx=10, pady=10)
        sb = tk.Scrollbar(lf, command=self.log.yview)
        self.log.configure(yscrollcommand=sb.set)
        self.log.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

        self.status = tk.StringVar(value="Ready")
        tk.Label(root, textvariable=self.status, bd=1, relief="sunken", anchor="w",
                 bg="#e5e7eb", font=("Segoe UI", 9)).pack(side="bottom", fill="x")

        miss = check_deps()
        if miss and not FROZEN:
            self.append("Missing: %s\nClick 'Install' once, then Run.\n\n" % ", ".join(miss))
        else:
            self.append("Ready. Set folders and press Run.\n\n")

    def append(self, s):
        self.log.insert(tk.END, s)
        self.log.see(tk.END)

    def logline(self, s):
        self.root.after(0, self.append, s)

    def set_status(self, s):
        self.root.after(0, self.status.set, s)

    def pick(self, var):
        d = filedialog.askdirectory()
        if d:
            var.set(d)

    def deps(self):
        def work():
            install_deps(self.logline)
            self.logline("\n")
        threading.Thread(target=work, daemon=True).start()

    def run(self):
        src = self.src.get().strip()
        if not os.path.isdir(src):
            messagebox.showerror("Error", "Folder not found:\n%s" % src)
            return
        if not FROZEN and check_deps():
            if not install_deps(self.logline):
                messagebox.showerror("Error", "Could not install dependencies.")
                return
        self.run_btn.config(state="disabled", text="⏳ Running...")
        self.set_status("Processing...")
        self.logline("\n" + "=" * 50 + "\nStarting...\n")

        def work():
            try:
                r = run_pipeline(src, self.out.get().strip(), self.logline)
                if r.get("status") == "success":
                    self.result = r["output_file"]
                    self.set_status("Done: %d units" % r["units"])
                    self.logline("\n✅ Done!  %d units\nFile: %s\n" % (r["units"], self.result))
                    self.root.after(0, lambda: messagebox.showinfo(
                        "Done", "%d units\n\nSaved to:\n%s" % (r["units"], self.result)))
                else:
                    self.set_status("Error")
                    self.logline("\n❌ %s\n" % r.get("message", "error"))
            except Exception as e:
                self.set_status("Error")
                self.logline("\n❌ %s\n" % e)
            finally:
                self.root.after(0, lambda: self.run_btn.config(state="normal", text="▶  Run"))
        threading.Thread(target=work, daemon=True).start()

    def open_result(self):
        p = self.result
        if (not p or not os.path.exists(p)):
            p = os.path.join(self.out.get().strip(), "Final_RealEstate_Database.xlsx")
        if not os.path.exists(p):
            messagebox.showinfo("No result", "Run the pipeline first.")
            return
        try:
            if os.name == "nt":
                os.startfile(p)
            elif sys.platform == "darwin":
                subprocess.call(["open", p])
            else:
                subprocess.call(["xdg-open", p])
        except Exception as e:
            messagebox.showerror("Error", str(e))

if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()